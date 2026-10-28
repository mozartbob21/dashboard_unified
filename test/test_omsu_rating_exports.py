"""Export parity and offline safety, with synthetic source data only."""
import base64
import hashlib
import io
import json
from html import unescape
from html.parser import HTMLParser

from openpyxl import load_workbook

from services.omsu_rating import calculate
from services.omsu_rating.export import workbook_bytes, snapshot_html, ZONES


def issue(name, fields, subject='Тестовая РСО'):
    return {'id': 1, 'assigned_to': {'name': name}, 'subject': subject,
            'custom_fields': [{'id': key, 'value': value} for key, value in {17: 1, **fields}.items()]}


def saved_rating():
    sources = {
        '4': {'data': [issue('ДУБНА', {70: 'Первая РСО', 165: 100, 166: 200, 169: 10000}),
                       issue('ИСТРА', {70: 'Вторая РСО', 165: 100, 166: 100, 169: 10000})]},
        '5': {'data': [issue('ДУБНА', {158: 1, 159: 4}), issue('ИСТРА', {158: 2, 159: 4}),
                       issue('КЛИН', {158: 3, 159: 10, 259: '337'}), issue('ВЛАСИХА', {158: 3, 159: 10, 259: '337'})]},
        '6': {'data': [issue('ДУБНА', {189: 20, 184: 100, 185: 10, 186: 5, 187: 3, 188: 2}),
                       issue('ИСТРА', {189: 2})]},
        '7': {'data': []},
    }
    for source in sources.values():
        source.update(updated_at='2026-10-28T10:00:00+00:00', error=None)
    return {**calculate.build(sources), 'updated_at': '2026-10-28T10:00:00+00:00',
            'refresh': {'status': 'success', 'error': None, 'finished_at': '2026-10-28T10:00:00+00:00'}, 'history': []}


def read_book(saved, **selection):
    return load_workbook(io.BytesIO(workbook_bytes(saved, **selection)))


def by_omsu(sheet):
    return {sheet.cell(row, 3).value: row for row in range(5, sheet.max_row + 1)}


def test_exact_original_headers_places_dynamics_and_all_nonparticipants():
    saved = saved_rating()
    book = read_book(saved)
    assert book.sheetnames == [f'Показатель {key}' for key in (4, 5, 6, 7)]
    assert [cell.value for cell in book['Показатель 4'][4]] == [
        'Место', 'Динамика (+ поднялся)', 'ОМСУ', 'Группа', 'Выполнение плана, %', 'Факт, % от НВВ', 'План, %', 'РСО']
    sheet = book['Показатель 5']
    indices = by_omsu(sheet)
    assert set(indices) == set(calculate.ALL_OMSU)
    assert sheet.cell(indices['ДУБНА'], 1).value == '1–2'
    assert sheet.cell(indices['ИСТРА'], 1).value == '1–2'
    assert sheet.cell(indices['ДУБНА'], 2).value == 2
    assert sheet.cell(indices['КЛИН'], 1).value == '51'
    assert sheet.cell(indices['КЛИН'], 4).value == 'обнулён'
    assert sheet.cell(indices['КЛИН'], 5).value == 0
    assert sheet.cell(indices['ВЛАСИХА'], 1).value == '5'
    assert sheet.cell(indices['ВЛАСИХА'], 4).value == 'ЗАТО, обнулён'
    assert sheet.cell(indices['БАЛАШИХА'], 1).value == '—'
    assert sheet.cell(indices['БАЛАШИХА'], 4).value == 'не участвует'
    assert sheet.cell(indices['ВОСХОД'], 4).value == 'ЗАТО, не участвует'
    assert indices['ВЛАСИХА'] > max(indices[name] for name in calculate.ALL_OMSU if name not in calculate.ZATO)


def test_selected_period_only_changes_its_indicator_and_rounding_is_display_only():
    saved = saved_rating()
    book = read_book(saved, indicator=5, period=0)
    assert '1 квартал' in book['Показатель 5']['A1'].value
    assert '1 полугодие' in book['Показатель 4']['A1'].value
    sheet = book['Показатель 5']
    assert sheet.cell(by_omsu(sheet)['ДУБНА'], 5).value == 1
    assert sheet.cell(by_omsu(sheet)['ДУБНА'], 2).value is None


def test_conditional_k6_and_preliminary_k7_keep_original_detail_columns():
    book = read_book(saved_rating())
    sheet = book['Показатель 6']
    assert [cell.value for cell in sheet[4]][5:7] == ['ЛС без ИПУ', 'Всего ЛС']
    row = by_omsu(sheet)['ДУБНА']
    assert sheet.cell(row, 1).value == '1–2'
    assert sheet.cell(row, 6).value == 20 and sheet.cell(row, 7).value == 100
    assert 'ЛС факт' in sheet['A3'].value and 'услов' in sheet['A3'].value.lower()
    assert '10 + 5 + 3 + 2' in sheet.cell(row, 6).comment.text
    sheet = book['Показатель 7']
    assert sheet['F4'].value == 'Мероприятий выполнено'
    row = by_omsu(sheet)['ДУБНА']
    assert sheet.cell(row, 6).value == '0 / 17'
    assert '0,00 (нет в таблице)' in sheet.cell(row, 7).value
    assert 'Предварительная' in sheet['A3'].value


def test_workbook_semantic_colors_and_formula_injection_protection():
    saved = saved_rating()
    item = saved['indicators']['5']
    malicious = '=HYPERLINK("https://invalid.test","x")'
    for rows in item['periods']:
        for row in rows:
            row['omsu'] = malicious if row['omsu'] == 'ДУБНА' else row['omsu']
            row['rso_names'] = ['=1+1', '+test', '-test', '@test']
    sheet = read_book(saved)['Показатель 5']
    row = by_omsu(sheet)[malicious]
    assert sheet.cell(row, 3).value == malicious and sheet.cell(row, 3).data_type == 's'
    assert sheet.cell(row, 6).value.startswith('=1+1') and sheet.cell(row, 6).data_type == 's'
    expected = next(r for r in item['periods'][item['default_period']] if r['omsu'] == malicious)
    assert sheet.cell(row, 3).fill.fgColor.rgb[-6:] == ZONES[expected['zone']][0]
    assert sheet.cell(row, 1).comment.text == expected['place_tooltip']
    assert sheet.cell(row, 5).comment.text == expected['value_tooltip']
    assert not any(cell.data_type == 'f' for row in sheet for cell in row)


class SnapshotParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.scripts = []
        self.policy = ''
        self.external = []
        self._script = None
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'script':
            self._script = {'attrs': attrs, 'text': ''}
        if tag == 'meta' and attrs.get('http-equiv') == 'Content-Security-Policy':
            self.policy = attrs['content']
        if attrs.get('src') and not attrs['src'].startswith('data:'):
            self.external.append(attrs['src'])
        if attrs.get('href'):
            self.external.append(attrs['href'])
    def handle_data(self, data):
        if self._script is not None:
            self._script['text'] += data
    def handle_endtag(self, tag):
        if tag == 'script' and self._script is not None:
            self.scripts.append(self._script)
            self._script = None


def test_snapshot_safe_interactive_offline_all_periods_and_excel_variants():
    saved = saved_rating()
    attack = '</script><img src=x onerror=alert(1)><script>alert(1)</script>'
    saved['indicators']['5']['periods'][0][0]['omsu'] = attack
    saved['unrelated_secret'] = 'secret-not-in-contract'
    html = snapshot_html(saved)
    parsed = SnapshotParser()
    parsed.feed(html)
    assert len(parsed.scripts) == 2 and not parsed.external
    assert "default-src 'none'" in parsed.policy and "connect-src 'none'" in parsed.policy
    executable = parsed.scripts[1]['text']
    digest = base64.b64encode(hashlib.sha256(executable.encode()).digest()).decode()
    assert f"script-src 'sha256-{digest}'" in parsed.policy
    assert 'unsafe-inline' not in parsed.policy.split('script-src ')[1].split(';')[0]
    assert parsed.scripts[0]['attrs']['id'] == 'or-snapshot-data'
    assert parsed.scripts[0]['attrs']['type'] == 'application/json'
    embedded = json.loads(parsed.scripts[0]['text'])
    assert embedded['indicators']['5']['periods'][0][0]['omsu'] == attack
    assert attack not in html and 'secret-not-in-contract' not in html
    assert '\\u003c/script\\u003e' in html
    for key in ('4:0', '4:1', '5:0', '5:1', '6:0', '7:2', 'default'):
        assert key in embedded['offline_exports']
    variant = load_workbook(io.BytesIO(base64.b64decode(embedded['offline_exports']['5:0'])))
    assert '1 квартал' in variant['Показатель 5']['A1'].value
    assert '1 полугодие' in variant['Показатель 4']['A1'].value
    assert 'or-detail' in html and 'or-search' in html and 'or-period' in html


def test_empty_export_has_explanation_instead_of_fabricated_ratings():
    saved = {**calculate.build({}), 'updated_at': None, 'refresh': {}, 'history': []}
    book = read_book(saved)
    assert book.sheetnames == ['Нет данных']
    assert 'Пока нет данных' in book.active['A2'].value
