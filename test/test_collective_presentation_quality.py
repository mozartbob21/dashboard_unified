"""Synthetic presentation regressions. No portal, account, model or external files."""
import io
import re
import unittest
from datetime import date
from unittest.mock import patch

from pptx import Presentation

from services.collective import presentation
from services.collective.client import HEADERS


def appeal(signatures='101', theme='ii', **fields):
    row = {key: '' for key in HEADERS}
    row.update({'ОМСУ': 'Тестовый округ', 'Статус': 'В работе',
                'Аннотация/краткое содержание': 'Просим восстановить водоснабжение.',
                'Ответ': 'Ремонт запланирован.', HEADERS[10]: signatures, '_theme': theme})
    row.update(fields)
    return row


def deck(rows):
    data = presentation.build(presentation.validate_rows(rows), date(2026, 1, 1), date(2026, 1, 2))
    return Presentation(io.BytesIO(data))


def detail_tables(prs):
    return [shape.table for slide in list(prs.slides)[2:-1]
            for shape in slide.shapes if shape.name == 'collective_rows']


def cells_text(tables):
    return '\n'.join(cell.text for table in tables for row in table.rows for cell in row.cells)


def effective_fonts(text_frame):
    """Only visible runs; paragraph defaults can supply the actual formatting."""
    return [(run.font, paragraph.font) for paragraph in text_frame.paragraphs
            for run in paragraph.runs if run.text.strip()]


class CollectivePresentationQualityTests(unittest.TestCase):
    def test_real_export_route_builds_filtered_editable_presentation(self):
        from test.test_collective import api
        with api({'modules': ['collective'], 'username': 'synthetic'}) as client:
            response = client.post('/mingkh/collective/api/pptx', json={
                'start': '2026-01-01', 'end': '2026-01-02',
                'rows': [appeal('101', **{'ОМСУ': 'Проверочный округ'})],
            })
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else '')
        self.assertEqual(response.headers['cache-control'], 'private, no-store')
        prs = Presentation(io.BytesIO(response.content))
        self.assertEqual(len(prs.slides), 4)
        self.assertEqual(re.sub(r'\s+', '', detail_tables(prs)[0].cell(1, 0).text), 'Проверочныйокруг')
        self.assertEqual(presentation.named(prs.slides[1], 'collective_totals').table.cell(2, 0).text, '1')

    def test_unknown_signatures_are_not_reported_as_zero_or_below_threshold(self):
        source = [appeal(''), appeal('Информация отсутствует'), appeal('0'), appeal('100'), appeal('101')]
        normalized = presentation.validate_rows(source)
        self.assertEqual([row['_sig'] for row in normalized], [None, None, 0, 100, 101])
        prs = deck(source)
        totals = presentation.named(prs.slides[1], 'collective_totals').table
        self.assertEqual([totals.cell(2, i).text for i in range(4)], ['5', '1', '2', '0'])
        unknown = presentation.named(prs.slides[1], 'collective_unknown_signatures').text
        self.assertRegex(unknown, r'\b2\b')
        self.assertRegex(unknown.lower(), r'неизвест|не указан|отсутств|нет данных')
        signatures = [row.cells[1].text for table in detail_tables(prs) for row in list(table.rows)[1:]]
        self.assertEqual(signatures.count('0'), 1)
        self.assertNotIn('None', signatures)

    def test_unknown_signatures_sort_after_known_zero(self):
        prs = deck([
            appeal('', **{'Аннотация/краткое содержание': 'UNKNOWN_APPEAL_MARKER'}),
            appeal('0', **{'Аннотация/краткое содержание': 'ZERO_APPEAL_MARKER'}),
            appeal('101', **{'Аннотация/краткое содержание': 'LARGE_APPEAL_MARKER'}),
        ])
        text = cells_text(detail_tables(prs))
        self.assertLess(text.index('LARGE_APPEAL_MARKER'), text.index('ZERO_APPEAL_MARKER'))
        self.assertLess(text.index('ZERO_APPEAL_MARKER'), text.index('UNKNOWN_APPEAL_MARKER'))

    def test_long_narratives_keep_all_tokens_in_editable_continuation_tables(self):
        annotation = ['ОПИСАНИЕ%04d' % i for i in range(450)] + ['КОНЕЦ_ОПИСАНИЯ']
        answer = ['РЕШЕНИЕ%04d' % i for i in range(700)] + ['КОНТРОЛЬНЫЙ_СРОК_15.10.2026']
        prs = deck([appeal(**{'Аннотация/краткое содержание': ' '.join(annotation),
                             'Ответ': ' '.join(answer)})])
        tables = detail_tables(prs)
        self.assertGreater(len(tables), 1)
        text = cells_text(tables)
        for token in annotation + answer:
            self.assertIn(token, text)
        totals = presentation.named(prs.slides[1], 'collective_totals').table
        self.assertEqual(totals.cell(2, 0).text, '1')  # continuations are not new appeals
        for table in tables:
            self.assertGreater(len(table.rows), 1)

    def test_mixed_topics_keep_signature_priority_without_extra_empty_pages(self):
        prs = deck([appeal('20', theme='ii'), appeal('200', theme='pr'), appeal('100', theme='kr')])
        tables = detail_tables(prs)
        self.assertEqual(len(tables), 1)
        self.assertEqual([r.cells[1].text for r in list(tables[0].rows)[1:]], ['200', '100', '20'])

    def test_numeric_validation_does_not_leak_internal_integer_errors(self):
        with self.assertRaisesRegex(ValueError, 'Проверьте исходные данные'):
            presentation.validate_rows([appeal('9' * 5000)])
        self.assertEqual(presentation.validate_rows([appeal('0' * 5000 + '42')])[0]['_sig'], 42)

    def test_comparison_signs_survive_html_cleanup(self):
        text = 'Давление < 3 атм, температура > 80°C; <b>проверить</b>.'
        self.assertEqual(presentation.clean(text), 'Давление < 3 атм, температура > 80°C; проверить.')
        self.assertEqual(presentation.clean('1 &lt; 2 &gt; 0'), '1 < 2 > 0')
        prs = deck([appeal(**{'Аннотация/краткое содержание': text})])
        self.assertIn('Давление < 3 атм, температура > 80°C', re.sub(r'\s+', ' ', cells_text(detail_tables(prs))))

    def test_each_content_slide_has_one_correct_page_number(self):
        prs = deck([appeal(theme='ii'), appeal(theme='kr'), appeal(theme='pr')])
        for index, slide in enumerate(prs.slides, 1):
            page_shapes = [shape for shape in slide.shapes if shape.name == 'collective_page']
            self.assertLessEqual(len(page_shapes), 1)
            if 1 < index < len(prs.slides):
                self.assertEqual(len(page_shapes), 1)
                self.assertEqual(page_shapes[0].text.strip(), str(index))
            for shape in slide.shapes:
                if shape.has_table:
                    self.assertLessEqual(shape.left + shape.width, prs.slide_width)
                    self.assertLessEqual(shape.top + shape.height, prs.slide_height)

    def test_body_is_readable_and_signatures_header_fits(self):
        prs = deck([appeal()])
        for table in detail_tables(prs):
            header = table.cell(0, 1)
            self.assertEqual(header.text, 'Подписей')
            inner_width_points = (table.columns[1].width - header.margin_left - header.margin_right) / 12700
            for font, default in effective_fonts(header.text_frame):
                size = font.size or default.size
                self.assertIsNotNone(size)
                # Arial Bold "Подписей" measures 58.90 pt at 12 pt.
                self.assertGreaterEqual(inner_width_points, 58.9 * size.pt / 12)
            for row in list(table.rows)[1:]:
                for cell in row.cells:
                    for font, default in effective_fonts(cell.text_frame):
                        size = font.size or default.size
                        self.assertIsNotNone(size)
                        self.assertGreaterEqual(size.pt, 13)

    def test_cover_period_is_readable_on_dark_background(self):
        prs = deck([appeal()])
        period = presentation.named(prs.slides[0], 'collective_period')
        self.assertIn('01.01.2026', period.text)
        self.assertIn('02.01.2026', period.text)
        for font, default in effective_fonts(period.text_frame):
            self.assertTrue(font.bold if font.bold is not None else default.bold)
            self.assertGreaterEqual((font.size or default.size).pt, 32)
            color = font.color if font.color.type is not None else default.color
            self.assertEqual(str(color.rgb), 'FFFFFF')

    def test_summary_title_has_contrast_on_original_green_band(self):
        prs = deck([appeal()])
        title = presentation.named(prs.slides[1], 'collective_title')
        self.assertIn('Коллективные', title.text)
        for font, default in effective_fonts(title.text_frame):
            self.assertTrue(font.bold if font.bold is not None else default.bold)
            self.assertGreaterEqual((font.size or default.size).pt, 24)
            color = font.color if font.color.type is not None else default.color
            self.assertEqual(str(color.rgb), 'FFFFFF')

    def test_presentation_limit_fails_with_actionable_russian_message(self):
        rows = presentation.validate_rows([
            appeal(**{'Аннотация/краткое содержание': ('Подробное описание %d. ' % i) * 400,
                      'Ответ': 'Подробный ответ ведомства. ' * 600}) for i in range(3)
        ])
        with patch.object(presentation, 'MAX_SLIDES', 5):
            with self.assertRaises(ValueError) as caught:
                presentation.build(rows, date(2026, 1, 1), date(2026, 1, 2))
        self.assertRegex(str(caught.exception).lower(), r'слайд|презентац')
        self.assertRegex(str(caught.exception).lower(), r'уменьш|уточн|сократ|до \d|не более')


if __name__ == '__main__':
    unittest.main()
