"""Conversion, isolation and outbound boundaries; no live AI or user data."""
import io
import json
import time
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

import pytest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image

from core.http_security import ToolsBodyLimit, private_data_path
from core.privacy import ai_endpoint, PrivacyError
from services.tools import documents, excel_merge, exe_builder, images, macros, pptx_converter, workspace
from services.tools.workspace import ToolError


def workbook(headers, rows):
    book=Workbook();sheet=book.active;sheet.append(headers)
    for row in rows:sheet.append(row)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
    output=io.BytesIO();book.save(output);return output.getvalue()


def picture():
    buffer=io.BytesIO();Image.new('RGB',(40,80),'navy').save(buffer,format='PNG');return buffer.getvalue()


def test_merge_aligns_headers_and_never_turns_text_into_formula(tmp_path):
    result=excel_merge.merge([('one.xlsx',workbook(['РСО','Остаток'],[['A',3]])),
       ('two.xlsx',workbook(['Остаток','РСО','Примечание'],[[5,'B','=HYPERLINK("https://example.invalid")']])),
       ('~$temp.xlsx',b'invalid'),(excel_merge.OUTPUT_NAME,b'invalid')],tmp_path)
    book=load_workbook(tmp_path/result['file'],data_only=False)
    assert result['rows']==2 and result['sources']==2
    assert list(book['Данные'].values)==[('РСО','Остаток','Примечание'),('A',3,None),('B',5,'=HYPERLINK("https://example.invalid")')]
    assert book['Данные']['C3'].data_type=='s'
    assert book['Источники'].max_row==3


def test_bad_excel_fails_without_partial_result(tmp_path):
    with pytest.raises(ToolError):excel_merge.merge([('good.xlsx',workbook(['A'],[[1]])),('bad.xlsx',b'bad')],tmp_path)
    assert not list(tmp_path.glob('*.xlsx'))


@pytest.mark.parametrize('member,content',[
 ('word/_rels/document.xml.rels','<Relationships><Relationship TargetMode="External" Type="http://x/image" Target="https://example.invalid/private"/></Relationships>'),
 ('word/document.xml','<doc xmlns:w="urn:w"><w:instrText>INCLUDETEXT "file:///private/secret"</w:instrText></doc>'),
 ('word/vbaProject.bin','untrusted'),
])
def test_office_external_content_rejected(member,content):
    data=io.BytesIO()
    with zipfile.ZipFile(data,'w') as z:z.writestr(member,content)
    with pytest.raises(ToolError):documents.validate_office(data.getvalue())


def test_uploaded_picture_embeds_without_losing_original_slide():
    html=images.embed_uploads('<h1>Original title</h1><p>Report</p>', [('photo.png',picture())])
    slides=pptx_converter.parse_html_to_slides_pro(html)
    assert any(s['title']=='Original title' for s in slides)
    assert any(s['images'] and s['images'][0].startswith('data:image/png') for s in slides)


def test_uploaded_picture_replaces_matching_html_reference():
    html=images.embed_uploads('<section><h2>Result</h2><img src="images/photo.png"></section>', [('photo.png',picture())])
    slides=pptx_converter.parse_html_to_slides_pro(html)
    assert len(slides)==1 and len(slides[0]['images'])==1
    assert 'images/photo.png' not in html


def test_ai_receives_no_image_payload_and_images_survive_missing_model_references():
    html=images.embed_uploads('<h1>Report</h1>', [('photo.png',picture())])
    with patch('services.summarizer.engine._qwen_chat',return_value='[{"title":"Report","texts":["Content"]}]') as ai:
        slides=pptx_converter.parse_html_to_slides_ai(html)
    assert 'base64' not in str(ai.call_args)
    assert any(s['images'] for s in slides)


@pytest.mark.parametrize('value',['http://aiplatform.mosreg.ru','https://evil.example','https://aiplatform.mosreg.ru.evil.example','https://user:pass@aiplatform.mosreg.ru','https://aiplatform.mosreg.ru/?key=secret','http://10.0.0.1'])
def test_unapproved_ai_is_rejected(value):
    with pytest.raises(PrivacyError):ai_endpoint(value)


@pytest.mark.parametrize('value',['https://aiplatform.mosreg.ru/api/user-models/v1','http://127.0.0.1:11434/v1','http://[::1]:8080/v1'])
def test_approved_ai(value):assert ai_endpoint(value)==value


def test_macro_download_not_executed(tmp_path):
    code='Option Explicit\nPublic Sub MakeReport()\nMsgBox "Ready"\nEnd Sub'
    with patch('services.summarizer.engine._qwen_chat',return_value=json.dumps({'code':code,'instructions':'Import module'})):
        result=macros.generate('Create a report on a new sheet','Excel VBA',tmp_path)
    assert result['code']==code and (tmp_path/'macro.bas').read_text(encoding='cp1251')==code
    with patch('services.summarizer.engine._qwen_chat',return_value=json.dumps({'code':code.replace('MsgBox','Shell')})):
        with pytest.raises(ToolError):macros.generate('Create a report on a new sheet','Excel VBA',tmp_path)


def test_exe_maps_only_job_and_separate_runtime(tmp_path):
    xml=ET.fromstring(exe_builder.sandbox_xml(tmp_path/'input',tmp_path/'output',tmp_path/'runtime'))
    assert xml.findtext('Networking')=='Disable'
    assert xml.findtext('ClipboardRedirection')=='Disable'
    mapped=xml.findall('MappedFolders/MappedFolder')
    assert [m.findtext('ReadOnly') for m in mapped]==['true','false','true']
    assert len(mapped)==3
    with patch.object(exe_builder,'available',return_value=False), patch.object(exe_builder.subprocess,'Popen') as spawn:
        with pytest.raises(ToolError):exe_builder.build('source.py',b'print("hi")',False,tmp_path)
        spawn.assert_not_called()


def test_pdf_to_docx_preserves_text(tmp_path):
    pymupdf=pytest.importorskip('pymupdf');pytest.importorskip('pdf2docx')
    pdf=pymupdf.open();page=pdf.new_page();page.insert_text((72,72),'Neurona conversion test')
    data=pdf.tobytes();pdf.close()
    result=documents.convert('pdf_docx','report.pdf',data,tmp_path)
    from docx import Document
    assert 'Neurona conversion test' in '\n'.join(p.text for p in Document(tmp_path/result['file']).paragraphs)


def test_word_to_pdf_has_been_removed(tmp_path):
    with pytest.raises(ToolError, match='только конвертация PDF'):
        documents.convert('docx_pdf', 'report.docx', b'not-used', tmp_path)
    assert not list(tmp_path.iterdir())


@pytest.fixture
def api(tmp_path,monkeypatch):
    monkeypatch.setattr(workspace,'ROOT',tmp_path/'accounts')
    from routers.tools import router
    app=FastAPI();app.include_router(router);app.add_middleware(ToolsBodyLimit,max_bytes=2*1024*1024)
    @app.middleware('http')
    async def user(request:Request,call_next):
        request.state.user={'id':int(request.headers.get('x-test-user','1')),'username':'test','role':'Пользователь','modules':[] if request.headers.get('x-no-grant') else ['tools']}
        return await call_next(request)
    with TestClient(app) as client:yield client


def test_job_download_is_private_and_exe_admin_only(api):
    r=api.post('/tools/api/run',data={'tool':'html_pptx','html':'<section><h1>Report</h1></section>'},files={'images':('photo.png',picture(),'image/png')})
    assert r.status_code==200,r.text
    job=r.json()['job_id']
    for _ in range(120):
        state=api.get('/tools/api/jobs/'+job).json()
        if state['state'] not in ('queued','running'):break
        time.sleep(.03)
    assert state['state']=='done',state
    result=api.get(state['url']);assert result.status_code==200
    with zipfile.ZipFile(io.BytesIO(result.content)) as z:assert any(n.startswith('ppt/media/') for n in z.namelist())
    assert api.get(state['url'],headers={'x-test-user':'2'}).status_code==404
    assert api.get('/tools/api/jobs/'+job,headers={'x-test-user':'2'}).status_code==404
    assert api.get('/tools/api/options',headers={'x-no-grant':'1'}).status_code==403
    assert api.post('/tools/api/run',data={'tool':'exe'}).status_code==403
    assert api.get('/tools/download/not-an-id/source.py').status_code==404


def test_upload_cannot_escape_workspace_and_request_is_bounded(api):
    assert api.post('/tools/template',files={'template_file':('../escape.pptx',b'bad')}).status_code==400
    assert api.post('/tools/api/run',content=b'x'*(2*1024*1024+1)).status_code==413


def test_private_raw_paths():
    for path in ['/data/tools/accounts/a/jobs/b/result.pptx','/data/auth/secrets.json','/data/cds/browser-profile/Cookies','/data/foo/token.json','/data/foo/private.pem','/data/foo/x.html']:
        assert private_data_path(path),path
    assert not private_data_path('/data/cds/appeals.xlsx')

@pytest.mark.parametrize('address',['127.0.0.1','10.10.34.2','169.254.169.254','192.168.1.1','::1','224.0.0.1'])
def test_remote_picture_cannot_access_private_network(address):
    from services.tools import external_images as remote
    with patch.object(remote.socket,'getaddrinfo',return_value=[(2,1,6,'',(address,443))]):
        with pytest.raises(ToolError):remote.public_address('test.example',443)


def test_remote_picture_embeds_and_failed_image_is_reported():
    from services.tools import external_images as remote
    with patch.object(remote,'download',side_effect=[picture(),ToolError('Unavailable')]) as download:
        html,failures=remote.embed_remote('<img src="https://example.invalid/a.png"><img src="https://example.invalid/b.png">')
    assert failures==1 and 'data:image/png;base64' in html
    assert download.call_count==2


def test_remote_redirect_is_rechecked_and_credentials_never_sent():
    from services.tools import external_images as remote
    class Reply:
        status=302
        def getheader(self,key):return 'http://127.0.0.1/private' if key=='Location' else None
    class Connection:
        def __init__(self,*args):self.sock=None
        def request(self,method,path,headers):
            assert set(headers)=={'Accept','User-Agent'}
        def getresponse(self):return Reply()
        def close(self):pass
    with patch.object(remote,'public_address',return_value='93.184.216.34') as resolve,patch.object(remote,'PinnedConnection',Connection):
        with pytest.raises(ToolError):remote.download('https://example.invalid/image.png')
        assert resolve.call_count==1  # HTTPS -> HTTP downgrade blocked before new request.
    with patch.object(remote.socket,'getaddrinfo',return_value=[(2,1,6,'',('127.0.0.1',80))]),patch.object(remote,'PinnedConnection') as connect:
        with pytest.raises(ToolError):remote.download('http://127.0.0.1/image.png')
        connect.assert_not_called()


def test_download_uses_validated_ip_even_if_dns_changes():
    from services.tools import external_images as remote
    with patch.object(remote.socket,'create_connection') as connect:
        conn=remote.PinnedConnection('public.example',80,'93.184.216.34',False,4)
        conn.connect()
        connect.assert_called_once_with(('93.184.216.34',80),4)


def test_local_ai_needs_no_external_api_key(monkeypatch):
    import httpx
    from services.summarizer.engine import _qwen_chat
    monkeypatch.setenv('QWEN_API_KEY','')
    response=httpx.Response(200,json={'choices':[{'message':{'content':'local answer'}}]},request=httpx.Request('POST','http://127.0.0.1:11434/v1/chat/completions'))
    with patch.object(httpx,'post',return_value=response) as post:
        assert _qwen_chat([{'role':'user','content':'synthetic'}],base='http://127.0.0.1:11434/v1')=='local answer'
        assert 'Authorization' not in post.call_args.kwargs['headers']
        assert post.call_args.kwargs['trust_env'] is False


def test_platform_context_never_reads_ungranted_modules():
    from services import platform_details as details
    with patch.object(details,'_load',return_value=[{'status':'critical','name':'allowed example'}]) as read:
        context=details.build_detailed_context('test',allowed_modules=['edo'])
        assert 'allowed example' in context
        read.assert_called_once_with(details.F_EDO)
    with patch.object(details,'_load') as read:
        details.build_detailed_context('test',allowed_modules=[])
        read.assert_not_called()


def test_removed_word_converter_is_not_exposed(api):
    assert 'docx_pdf' not in api.get('/tools/api/options').json()
    response = api.post('/tools/api/run', data={'tool': 'docx_pdf'}, files={'file': ('test.docx', b'not-used')})
    assert response.status_code == 400


def test_coordinate_upload_download_and_account_isolation(api):
    data = workbook(['Широта', 'Долгота', 'Заявка'], [[55.7, 37.5, 'A'], [55.701, 37.5, 'B']])
    response = api.post('/tools/api/run', data={'tool': 'coordinates', 'radius': '500', 'prefix': 'VZ'},
                        files={'file': ('points.xlsx', data), 'registry': ('', b'')})
    assert response.status_code == 200, response.text
    job = response.json()['job_id']
    for _ in range(120):
        state = api.get('/tools/api/jobs/' + job).json()
        if state['state'] not in ('queued', 'running'): break
        time.sleep(.03)
    assert state['state'] == 'done', state
    result = api.get(state['url'])
    assert result.status_code == 200
    book = load_workbook(io.BytesIO(result.content))
    assert book['Объекты']['A2'].value == 'VZ-0001'
    assert book['Записи'].max_row == 3
    assert api.get(state['url'], headers={'x-test-user': '2'}).status_code == 404
    assert api.get('/tools/api/jobs/' + job, headers={'x-test-user': '2'}).status_code == 404


def finished_job(api, response):
    assert response.status_code == 200, response.text
    job = response.json()['job_id']
    for _ in range(150):
        state = api.get('/tools/api/jobs/' + job).json()
        if state['state'] not in ('queued', 'running'): return state
        time.sleep(.03)
    raise AssertionError('Job did not finish')


def test_pdf_pages_api_preserves_order_and_private_results(api):
    from test.test_pdf_pages import document, texts
    files = [('file', ('one.pdf', document(['One']))), ('file', ('two.pdf', document(['Two'])))]
    assert api.post('/tools/api/pdf-info', files=files).json()['files'] == [
        {'name':'one.pdf','pages':1}, {'name':'two.pdf','pages':1}]
    assert api.post('/tools/api/pdf-info', files=files, headers={'x-no-grant':'1'}).status_code == 403
    state = finished_job(api, api.post('/tools/api/run', data={'tool':'pdf_pages','action':'merge'}, files=files[::-1]))
    assert state['state'] == 'done', state
    assert texts(api.get(state['url']).content) == ['Two','One']
    assert api.get(state['url'], headers={'x-test-user':'2'}).status_code == 404
    state = finished_job(api, api.post('/tools/api/run', data={'tool':'pdf_pages','action':'extract','pages':'2'}, files=files[:1]))
    assert state['state'] == 'error' and 'url' not in state


def test_branding_is_shared_but_presentation_output_is_private(api):
    from test.test_tools_branding import template
    from services.tools import branding
    other = {'x-test-user':'2'}
    response = api.post('/tools/template', files={'template_file':('Common.pptx',template())})
    assert response.status_code == 200, response.text
    assert 'Common.pptx' in api.get('/tools/api/options', headers=other).json()['templates']
    assert api.post('/tools/emblem', files={'emblem_file':('organization.png',picture())}).status_code == 200
    emblem = api.get('/tools/emblem', headers=other)
    assert emblem.content == branding.emblem_png(picture())
    assert api.get('/tools/emblem', headers={'x-no-grant':'1'}).status_code == 403
    for name in ['', 'Common.pptx']:
        response = api.post('/tools/api/run', headers=other,
                            data={'tool':'html_pptx','html':'<h1>Shared branding</h1>','template':name})
        assert response.status_code == 200, response.text
        job = response.json()['job_id']
        for _ in range(150):
            state = api.get('/tools/api/jobs/'+job, headers=other).json()
            if state['state'] not in ('queued','running'): break
            time.sleep(.03)
        assert state['state'] == 'done', state
        with zipfile.ZipFile(io.BytesIO(api.get(state['url'], headers=other).content)) as z:
            assert emblem.content in [z.read(n) for n in z.namelist() if n.startswith('ppt/media/')]
        assert api.get(state['url']).status_code == 404
    deleted = api.post('/tools/template/delete', headers=other, data={'name':'Common.pptx'})
    assert deleted.status_code == 200
    assert 'Common.pptx' not in api.get('/tools/api/options').json()['templates']
