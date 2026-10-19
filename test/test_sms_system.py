"""Offline synthetic contracts for the imported system-address methodology."""
import base64
import datetime as dt
import importlib.util
import io
from pathlib import Path
import sys
import types
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook


@pytest.fixture
def modules(tmp_path, monkeypatch):
    source = Path(__file__).resolve().parents[1] / 'services' / 'mingkh' / 'sms'
    package = types.ModuleType('sms_system_unit')
    package.__path__ = [str(source)]
    monkeypatch.setitem(sys.modules, package.__name__, package)
    context = types.ModuleType(package.__name__ + '.context')
    context.DATA = tmp_path
    context.today = dt.date.today
    client = types.ModuleType(package.__name__ + '.client')
    client.Pentaho = lambda: (_ for _ in ()).throw(AssertionError('No external requests in tests'))
    dataset = types.ModuleType(package.__name__ + '.dataset')
    dataset.WORKERS = 1
    dataset.chunks = lambda start, end: [(start, end)]
    for module in (context, client, dataset):
        monkeypatch.setitem(sys.modules, module.__name__, module)
    loaded = []
    for name in ('system_engine', 'sys_data'):
        spec = importlib.util.spec_from_file_location(package.__name__ + '.' + name, source / (name + '.py'))
        module = importlib.util.module_from_spec(spec)
        monkeypatch.setitem(sys.modules, module.__name__, module)
        spec.loader.exec_module(module)
        setattr(package, name, module)
        loaded.append(module)
    return loaded


PROD = 'Жалобы продолжаются'
NET = 'Нет жалоб последние 30 дней'
STREET = 'Город А, улица Старая'
NEW = 'Город А, улица Новая'


def registry(status=PROD, street=STREET, omsu='Округ А', rows=None, formula=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sheet1'
    sheet.append(['Улица', 'ОМСУ', 'Статус', 'Номера задач #'])
    for row in rows or [[street, omsu, status, '=HYPERLINK("https://invalid")' if formula else '123']]:
        sheet.append(row)
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def complaints_file(path, records):
    cols = ['Номер ЕЦУР', 'Номер в источнике', 'Дата (первого взятия в работу)', 'Тип сообщения', 'Адрес', 'Дом', 'ОМСУ']
    pd.DataFrame(records, columns=cols).to_excel(path, index=False, sheet_name='Sheet1')
    return path


def complaint(day, street=NEW, idx=1, typ=0, omsu='Округ А', house='1'):
    return [idx, f'C-{idx}', day, typ, street, house, omsu]


def process_case(tmp_path, modules, records, ref='2025-07-01', status=PROD):
    engine, service = modules
    r = tmp_path / 'registry.xlsx'
    r.write_bytes(registry(status))
    c = complaints_file(tmp_path / 'complaints.xlsx', records)
    return engine.process(r, c, pd.Timestamp(ref))


@pytest.mark.parametrize('dates,expected', [
    (['2025-06-10', '2025-06-10', '2025-06-11'], 1),
    (['2025-06-10', '2025-06-10', '2025-06-10'], 0),
    (['2025-06-10', '2025-06-11'], 0),
    (['2025-06-01', '2025-06-10', '2025-06-11'], 0),
    (['2025-06-02', '2025-06-10', '2025-06-11'], 1),
])
def test_new_street_requires_three_complaints_two_days_30_day_window(tmp_path, modules, dates, expected):
    _, new, *_ = process_case(tmp_path, modules, [complaint(day, idx=i) for i, day in enumerate(dates)])
    assert len(new) == expected


def test_year_from_reference_last_day_included_future_excluded(tmp_path, modules):
    engine, _ = modules
    path = complaints_file(tmp_path / 'date-boundary.xlsx', [
        complaint('2025-07-01 23:59:59', idx=1),
        complaint('2025-07-02 00:00:00', idx=2),
        complaint('2024-07-01', idx=3),
        complaint('2026-01-01', idx=4),
        complaint('2025-07-01', idx=5, typ=1),
    ])
    result = engine.load_complaints(path, pd.Timestamp('2025-07-01'))
    assert result['номер'].tolist() == ['C-1']


def test_empty_complaints_still_yield_old_street_as_no_complaints(tmp_path, modules):
    _, _, streets, *_ = process_case(tmp_path, modules, [])
    assert streets['Статус'].tolist() == [NET]


@pytest.mark.parametrize('houses,expected', [(['1', '1'], PROD), (['1', '2'], NET), (['', ''], PROD)])
def test_returned_street_requires_two_days_at_same_house_or_no_house(tmp_path, modules, houses, expected):
    records = [complaint('2025-06-10', STREET, 1, house=houses[0]), complaint('2025-06-11', STREET, 2, house=houses[1])]
    _, _, streets, *_ = process_case(tmp_path, modules, records, status=NET)
    assert streets['Статус'].tolist() == [expected]


def test_existing_active_street_needs_one_recent_complaint(tmp_path, modules):
    _, _, streets, *_ = process_case(tmp_path, modules, [complaint('2025-07-01 18:00', STREET)], status=PROD)
    assert streets['Статус'].tolist() == [PROD]


def test_import_registry_and_previous_survive_reload(modules):
    _, service = modules
    info = service.import_registry(registry(), '2025-06-20')
    assert info['streets'] == 1
    assert service.registries() == [dt.date(2025, 6, 20)]
    assert service.previous(dt.date(2025, 7, 1)) == dt.date(2025, 6, 20)
    with pytest.raises(ValueError, match='раньше'):
        service.previous(dt.date(2025, 6, 20))


@pytest.mark.parametrize('payload', [b'', b'not Excel', b'PK malformed workbook'])
def test_bad_upload_is_rejected_without_replacing_previous(modules, payload):
    _, service = modules
    service.import_registry(registry(), '2025-06-20')
    path = service.REG_DIR / '2025-06-20.xlsx'
    saved = path.read_bytes()
    with pytest.raises(ValueError):
        service.import_registry(payload, '2025-06-20')
    assert path.read_bytes() == saved


@pytest.mark.parametrize('bad_rows', [
    [[STREET, '', PROD, '123']],
    [[STREET, 'Округ А', 'Неизвестно', '123']],
    [[STREET, 'Округ А', PROD, '123'], [STREET, 'Округ Б', PROD, '124']],
])
def test_required_registry_information_and_ambiguous_street_duplicates(modules, bad_rows):
    _, service = modules
    with pytest.raises(ValueError):
        service.import_registry(registry(rows=bad_rows), '2025-06-20')
    assert service.registries() == []


def test_upload_requires_headers_and_rejects_formula(modules):
    _, service = modules
    with pytest.raises(ValueError, match='формул'):
        service.import_registry(registry(formula=True), '2025-06-20')
    workbook = Workbook()
    workbook.active.append(['Адрес'])
    workbook.active.append(['Улица1'])
    b = io.BytesIO(); workbook.save(b)
    with pytest.raises(ValueError, match='колонки'):
        service.import_registry(b.getvalue(), '2025-06-20')


def test_upload_size_and_archive_limits(modules, monkeypatch):
    _, service = modules
    payload = registry()
    monkeypatch.setattr(service, 'MAX_REGISTRY_BYTES', len(payload) - 1)
    with pytest.raises(ValueError, match='20 МБ'):
        service.import_registry(payload, '2025-06-20')
    monkeypatch.setattr(service, 'MAX_REGISTRY_BYTES', len(payload) + 1)
    monkeypatch.setattr(service, 'MAX_EXPANDED_BYTES', 10)
    with pytest.raises(ValueError, match='объём'):
        service.import_registry(payload, '2025-06-20')


def test_registry_dates_are_not_paths(modules):
    _, service = modules
    for value in ('../../anything', '2025-99-99', '2999-01-01', '20250620', None):
        with pytest.raises(ValueError):
            service.import_registry(registry(), value)


def test_run_commits_after_complete_exports_and_cleans_workdir(tmp_path, modules, monkeypatch):
    engine, service = modules
    service.import_registry(registry(), '2025-06-20')
    def fetch(date_to, path, log):
        complaints_file(path, [complaint('2025-06-22', idx=1), complaint('2025-06-23', idx=2), complaint('2025-07-01 19:00', idx=3)])
        return 3
    monkeypatch.setattr(service, 'complaints', fetch)
    output = service.run({'date': '2025-07-01'})
    assert output['counts']['new'] == 1
    assert output['counts']['total'] == 2
    assert len(service.registries()) == 2
    assert list(service.OUT.iterdir()) == []
    data = base64.b64decode(output['files']['registry']['b64'])
    assert data == (service.REG_DIR / '2025-07-01.xlsx').read_bytes()
    book = load_workbook(io.BytesIO(data))
    assert book['Sheet1']['C1'].value == 'Жалоб за 2025'
    assert book['Лист1']['D1'].value == 'Жалоб за 2025'


def test_failed_export_never_overwrites_registry(tmp_path, modules, monkeypatch):
    engine, service = modules
    service.import_registry(registry(), '2025-06-20')
    service.import_registry(registry(NET), '2025-07-01')
    saved = (service.REG_DIR / '2025-07-01.xlsx').read_bytes()
    monkeypatch.setattr(service, 'complaints', lambda date_to, path, log: (complaints_file(path, []), 0)[1])
    monkeypatch.setattr(engine, 'write_returned_and_new', lambda *a: (_ for _ in ()).throw(OSError('Synthetic export failure')))
    with pytest.raises(OSError):
        service.run({'date': '2025-07-01'})
    assert saved == (service.REG_DIR / '2025-07-01.xlsx').read_bytes()
    assert list(service.OUT.iterdir()) == []


def test_exported_strings_are_never_formulas(tmp_path, modules):
    engine, _ = modules
    adr, nw, g, s_str, _ = process_case(tmp_path, modules, [complaint('2025-06-10', STREET)])
    adr['задачи'] = adr['задачи'].astype(object)
    adr.loc[0, 'задачи'] = '=HYPERLINK("https://invalid")'
    g.loc[0, 'задачи'] = '=HYPERLINK("https://invalid")'
    target = tmp_path / 'export.xlsx'
    engine.write_excel(adr, nw, g, s_str, pd.Timestamp('2025-07-01'), target)
    workbook = load_workbook(target, data_only=False)
    assert workbook['Sheet1']['D2'].data_type == 's'
    assert workbook['Sheet1']['D2'].value.startswith('=')
    literal = [row[6] for row in workbook['Лист1'].iter_rows(min_row=2) if str(row[6].value).startswith('=')]
    assert literal and all(cell.data_type == 's' for cell in literal)


def test_january_window_includes_december_but_year_count_stays_separate(tmp_path, modules):
    engine, _ = modules
    records = [complaint('2024-12-20', idx=1), complaint('2024-12-31', idx=2), complaint('2025-01-02 20:15', idx=3)]
    adr, new, streets, source, stats = process_case(tmp_path, modules, records, ref='2025-01-02')
    assert len(new) == 1
    assert int(new.iloc[0]['кол']) == 3
    assert int(new.iloc[0]['кол_год']) == 1
    target = tmp_path / 'january.xlsx'
    engine.write_excel(adr, new, streets, source, pd.Timestamp('2025-01-02'), target)
    book = load_workbook(target)
    new_row = next(row for row in book['Sheet1'].iter_rows(values_only=True) if row[0] == NEW)
    assert new_row[2] == 1
    assert book['Новые системные улицы']['D2'].value == 3


def test_prepare_does_not_save_and_upload_is_idempotent_but_never_overwrites(modules):
    _, service = modules
    assert service.prepare_registry(registry()).startswith(b'PK')
    assert service.registries() == []
    assert service.import_registry(registry(), '2025-06-20')['created'] is True
    before = (service.REG_DIR / '2025-06-20.xlsx').read_bytes()
    assert service.import_registry(registry(), '2025-06-20')['created'] is False
    with pytest.raises(ValueError, match='другой реестр'):
        service.import_registry(registry(NET), '2025-06-20')
    assert before == (service.REG_DIR / '2025-06-20.xlsx').read_bytes()


def test_invalid_portal_data_is_not_interpreted_as_zero(tmp_path, modules):
    engine, _ = modules
    bad = tmp_path / 'missing.xlsx'
    pd.DataFrame({'Дата': ['2025-06-01'], 'Other': ['something']}).to_excel(bad, index=False)
    with pytest.raises(ValueError):
        engine.load_complaints(bad, pd.Timestamp('2025-07-01'))
    bad = complaints_file(tmp_path / 'invalid-date.xlsx', [complaint('not-a-date')])
    with pytest.raises(ValueError, match='распознать даты'):
        engine.load_complaints(bad, pd.Timestamp('2025-07-01'))


def test_same_street_text_does_not_count_complaints_from_other_municipality(tmp_path, modules):
    _, _, streets, *_ = process_case(tmp_path, modules, [complaint('2025-06-10', STREET, omsu='Округ Б')], status=NET)
    assert streets['кол'].tolist() == [0]
    assert streets['Статус'].tolist() == [NET]


def test_indistinguishable_street_names_in_different_municipalities_are_rejected(tmp_path, modules):
    records = [complaint('2025-06-10', STREET, idx=1, omsu='Округ Б'),
               complaint('2025-06-11', STREET, idx=2, omsu='Округ Б'),
               complaint('2025-06-12', STREET, idx=3, omsu='Округ Б')]
    with pytest.raises(ValueError, match='разным ОМСУ'):
        process_case(tmp_path, modules, records)


def test_portal_january_count_is_year_only_but_workbook_keeps_december(tmp_path, modules, monkeypatch):
    engine, service = modules
    columns = ['Внутренний Id', 'Номер ЕЦУР', 'Номер в источнике', 'Дата (первого взятия в работу)', 'Тип сообщения', 'Адрес', 'Дом', 'ОМСУ']
    rows = [[index, *complaint(day, idx=index)] for index, day in enumerate(
        ('2024-12-20', '2024-12-31', '2025-01-02 23:59:59'), 1)]
    calls = []
    class FakePortal:
        def periods(self, *args, **kwargs):
            return kwargs
        def query(self, name, params, timeout):
            calls.append((name, params))
            return columns, rows
    monkeypatch.setattr(service, 'Pentaho', FakePortal)
    monkeypatch.setattr(service, 'today', lambda: dt.date(2025, 1, 3))
    messages = []
    path = tmp_path / 'portal-complaints.xlsx'
    count = service.complaints(dt.date(2025, 1, 2), path, messages.append)
    assert count == 1
    assert len(pd.read_excel(path)) == 3
    assert calls[0][1]['curr_period_start'] == '30 day'
    assert calls[0][1]['curr_period_end'] == '1 day'
    assert messages == ['Жалоб ВС/ВО за 2025 год по дату отчёта: 1']
    parsed = engine.load_complaints(path, pd.Timestamp('2025-01-02'))
    assert len(parsed) == 3 and int(parsed['отчётный_год'].sum()) == 1
