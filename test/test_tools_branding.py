import io
from pathlib import Path
from unittest.mock import patch

import pytest
from PIL import Image
from pptx import Presentation
from services.tools import branding, workspace, pptx_converter as pptx
from services.tools.workspace import ToolError


def template(title='Organization'):
    prs=Presentation(); prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text=title
    result=io.BytesIO(); prs.save(result); return result.getvalue()


def image(color):
    result=io.BytesIO(); Image.new('RGB',(30,40),color).save(result,format='PNG'); return result.getvalue()


@pytest.fixture
def storage(tmp_path, monkeypatch):
    monkeypatch.setattr(workspace,'ROOT',tmp_path/'accounts')
    return tmp_path


def test_migration_preserves_conflicting_templates_and_originals(storage):
    old=storage/'templates'; old.mkdir(); (old/'shared.pptx').write_bytes(template('Old common'))
    personal=storage/'accounts'/'one'/'templates'; personal.mkdir(parents=True)
    (personal/'shared.pptx').write_bytes(template('Previously personal'))
    (personal/'damaged.pptx').write_bytes(b'not PowerPoint')
    before=(personal/'shared.pptx').read_bytes()
    data=branding.options()
    assert len(data['templates'])==2
    assert 'shared.pptx' in data['templates']
    assert len(data['branding_warnings'])==1
    assert (personal/'shared.pptx').read_bytes()==before
    assert (personal/'damaged.pptx').exists()
    branding.delete_template('shared.pptx')
    assert 'shared.pptx' not in branding.options()['templates']  # Not reimported on next visit.


def test_legacy_common_emblem_preferred_and_alternatives_preserved(storage):
    (storage/'emblem.png').write_bytes(image('navy'))
    account=storage/'accounts'/'one';account.mkdir(parents=True)
    (account/'emblem.png').write_bytes(image('red'))
    options=branding.options()
    assert options['emblem_url']
    assert branding.capture('')['emblem.png']==branding.emblem_png(image('navy'))
    assert len(list((branding.root()/'previous-emblems').glob('*.png')))==2
    assert options['branding_warnings']


def test_running_job_uses_snapshot_despite_later_common_changes(storage):
    branding.save_template('Brand.PPTX',template('Before'))
    branding.save_emblem(image('navy'))
    assets=branding.capture('Brand.PPTX')
    branding.save_template('Brand.PPTX',template('After'))
    branding.save_emblem(image('red'))
    job=storage/'job';job.mkdir()
    with pptx.workspace(branding.prepare_job(job,assets)):
        assert pptx._load_emblem()==branding.emblem_png(image('navy'))
        assert Presentation(str(pptx.templates_dir()/'Brand.PPTX')).slides[0].shapes.title.text=='Before'
    assert branding.capture('Brand.PPTX')!=assets
    assert 'Brand.PPTX' in branding.options()['templates']


def test_shared_emblem_used_in_standard_presentation(storage):
    branding.save_emblem(image('navy'))
    job=storage/'job';job.mkdir()
    with pptx.workspace(branding.prepare_job(job,branding.capture(''))):
        output=pptx.build_pptx(pptx.parse_html_to_slides_pro('<h1>Report</h1>'),None,str(job/'result.pptx'))
    prs=Presentation(output)
    assert any(getattr(shape,'image',None) for shape in prs.slides[0].shapes)


def test_invalid_assets_do_not_replace_common_design(storage):
    branding.save_template('Brand.pptx',template());branding.save_emblem(image('navy'))
    before=branding.capture('Brand.pptx')
    with pytest.raises(ToolError):branding.save_template('Brand.pptx',b'bad')
    with pytest.raises(ToolError):branding.save_emblem(b'bad')
    with pytest.raises(ToolError):branding.capture('../secret.pptx')
    assert branding.capture('Brand.pptx')==before
