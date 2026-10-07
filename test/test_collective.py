"""Synthetic contract/access/export tests. Never contact the live portal."""
import io
import ssl
import unittest
from datetime import date
from pathlib import Path
from urllib.parse import parse_qs
from unittest.mock import patch
from zipfile import ZipFile

import httpx
from fastapi import FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.testclient import TestClient
from pptx import Presentation
from routers import collective
from services.collective import client, presentation


def record(theme='ii', signatures='101', **fields):
    row = {key: '' for key in client.HEADERS}
    row.update({'ОМСУ': 'Тестовый округ', 'Дата поступления обращения': '06.10.2026',
                'Аннотация/краткое содержание': 'Просим восстановить водоснабжение.',
                'Статус': 'В работе', 'Ответ': 'Ремонт запланирован.',
                client.HEADERS[10]: signatures, '_theme': theme, '_sig': 99999})
    row.update(fields)
    return row


def api(user):
    app = FastAPI()
    @app.middleware('http')
    async def add_user(request, call_next):
        request.state.user = user
        return await call_next(request)
    app.include_router(collective.router)
    app.mount('/static', StaticFiles(directory=Path(__file__).resolve().parents[1] / 'static'), name='static')
    return TestClient(app)


class CollectiveTests(unittest.TestCase):
    def test_protocol_basic_auth_verified_tls_and_repeated_filters(self):
        seen = []
        original = httpx.Client
        def handler(request):
            seen.append(request)
            return httpx.Response(200, json={'resultset': [[None] * 16]})
        def transport(**kwargs):
            self.assertIsInstance(kwargs['verify'], ssl.SSLContext)
            self.assertEqual(kwargs['verify'].verify_mode, ssl.CERT_REQUIRED)
            self.assertFalse(kwargs['follow_redirects'])
            self.assertFalse(kwargs['trust_env'])
            return original(transport=httpx.MockTransport(handler), **kwargs)
        params = client.query_params(date(2026, 10, 1), date(2026, 10, 7), [('17', 'A'), ('22', 'B')], date(2026, 10, 7))
        with patch.object(client.httpx, 'Client', side_effect=transport):
            client.CollectivePortal('synthetic', 'not-a-real-password').query('q_download', params)
        self.assertEqual(len(seen), 1)
        self.assertEqual(seen[0].method, 'POST')
        self.assertEqual(seen[0].url.host, 'cur.bi.mosreg.ru')
        self.assertNotIn('password', str(seen[0].url))
        body = parse_qs(seen[0].content.decode())
        self.assertEqual(body['path'], ['/public/jalobi/jalobi.cda'])
        self.assertEqual(body['paramvedom_kurat'], ['17', '22'])
        self.assertEqual(body['paramcurr_period_start'], ['6 day'])
        self.assertEqual(body['paramall_per'], ['0'])
        with self.assertRaises(client.PortalError):
            client.CollectivePortal('x', 'y').query('arbitrary')

    def test_redirect_and_invalid_payload_cannot_leak_or_look_like_empty_data(self):
        original = httpx.Client
        for response in [httpx.Response(302, headers={'location': 'https://example.invalid/collect'}),
                         httpx.Response(200, json={'resultset': [['wrong', 'columns']]}),
                         httpx.Response(200, json={'resultset': {'bad': 'shape'}})]:
            seen = []
            def handler(request):
                seen.append(request)
                return response
            with patch.object(client.httpx, 'Client', side_effect=lambda **kw: original(transport=httpx.MockTransport(handler), **kw)):
                with self.assertRaises(client.PortalError):
                    client.CollectivePortal('x', 'secret').query('q_download')
            self.assertEqual(len(seen), 1)

    def test_curators_include_combined_names_and_never_send_empty_filter(self):
        portal = client.CollectivePortal('x', 'y')
        name = client.CURATOR.upper()
        with patch.object(portal, 'query', side_effect=[[[1, name], [2, name + ' + объединение'], [3, 'Другое ведомство']], [[''] * 16]]) as query:
            kurators, rows = portal.fetch(date(2026, 10, 1), date(2026, 10, 7))
        self.assertEqual([x[0] for x in kurators], ['1', '2'])
        self.assertEqual(query.call_args.args[1]['vedom_kurat'], ['1', '2'])
        with patch.object(portal, 'query', return_value=[]) as query:
            with self.assertRaisesRegex(client.PortalError, 'найти куратора'):
                portal.fetch(date.today(), date.today())
            self.assertEqual(query.call_count, 1)

    def test_invalid_matching_curator_ids_never_fetch_all_departments(self):
        portal = client.CollectivePortal('x', 'y')
        for invalid in [None, '', '  ', '-1', -1]:
            with self.subTest(invalid=invalid), patch.object(portal, 'query', return_value=[[17, client.CURATOR], [invalid, client.CURATOR]]) as query:
                with self.assertRaisesRegex(client.PortalError, 'идентификатор'):
                    portal.fetch(date.today(), date.today())
                self.assertEqual(query.call_count, 1)

    def test_periods_and_today_contract(self):
        today = date(2026, 10, 7)
        self.assertEqual(client.query_params(today, today, [], today)['all_per'], '1')
        for start, end in [('bad', '2026-10-07'), ('2026-10-08', '2026-10-07'), ('2026-10-07', '2026-10-08')]:
            with self.assertRaises(ValueError):
                client.period(start, end, today)

    def test_access_control_and_shared_admin_credentials(self):
        for user in [None, {'modules': []}]:
            c = api(user)
            for path in ['/mingkh/collective', '/mingkh/collective/api/data']:
                self.assertEqual(c.get(path).status_code, 403)
            self.assertEqual(c.post('/mingkh/collective/api/pptx', json={}).status_code, 403)
        c = api({'modules': ['mingkh'], 'username': 'synthetic'})
        with patch.object(collective, 'credentials', return_value=None):
            response = c.get('/mingkh/collective/api/data?start=2026-01-01&end=2026-01-02')
            self.assertEqual(response.status_code, 503)
            self.assertIn('Администратору', response.json()['error'])
        with patch.object(collective, 'credentials', return_value={'username': 'u', 'password': 'p'}) as credentials, \
             patch.object(collective.CollectivePortal, 'fetch', return_value=([('17', 'Куратор')], [[''] * 16])):
            response = c.get('/mingkh/collective/api/data?start=2026-01-01&end=2026-01-02')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()['start'], '2026-01-01')
            self.assertEqual(response.json()['head'], client.HEADERS)
            self.assertEqual(response.headers['cache-control'], 'private, no-store')
            credentials.assert_called_once_with('mingkh')

    def test_export_validation_and_manual_themes(self):
        rows = presentation.validate_rows([record('kr', '100'), record('pr', '101')])
        self.assertEqual(rows[0]['_theme'], 'kr')
        self.assertEqual(rows[0]['_sig'], 100)
        self.assertEqual(rows[1]['_sig'], 101)
        for bad in [None, [], [record('__proto__')], [record(**{'Ответ': []})]]:
            with self.assertRaises(ValueError):
                presentation.validate_rows(bad)
        c = api({'modules': ['mingkh']})
        self.assertEqual(c.post('/mingkh/collective/api/pptx', json={'start': '2026-01-01', 'end': '2026-01-02', 'rows': []}).status_code, 400)
        with patch.object(collective, 'MAX_EXPORT_BYTES', 1):
            self.assertEqual(c.post('/mingkh/collective/api/pptx', content=b'{}').status_code, 413)

    def test_presentation_has_filtered_data_counts_and_editable_tables(self):
        rows = presentation.validate_rows([
            record('kr', '100', **{'Статус': 'Не учитывается'}),
            record('pr', '101', **{'Повтор': 'Да', 'ОМСУ': 'Другой тестовый округ'}),
        ])
        data = presentation.build(rows, date(2026, 1, 1), date(2026, 1, 2))
        with ZipFile(io.BytesIO(data)) as archive:
            self.assertFalse(any('thumbnail' in name.lower() for name in archive.namelist()))
            self.assertNotIn(b'thumbnail', archive.read('_rels/.rels'))
        prs = Presentation(io.BytesIO(data))
        self.assertEqual(len(prs.slides), 5)  # cover, summary, two theme slides, closing
        summary = prs.slides[1]
        totals = presentation.named(summary, 'collective_totals').table
        self.assertEqual([totals.cell(2, i).text for i in range(4)], ['2', '1', '1', '1'])
        themes = presentation.named(summary, 'collective_themes').table
        self.assertEqual(themes.cell(3, 1).text, '0')  # excluded KR appeal
        self.assertEqual(themes.cell(4, 1).text, '1')
        table = presentation.named(prs.slides[2], 'collective_rows').table
        self.assertEqual(table.cell(1, 1).text, '100')
        self.assertIn('Не учитывается', table.cell(1, 3).text)
        all_text = '\n'.join(shape.text if shape.has_text_frame else '\n'.join(c.text for r in shape.table.rows for c in r.cells) if shape.has_table else '' for slide in prs.slides for shape in slide.shapes)
        self.assertIn('01.01.2026', all_text)
        self.assertIn('Прочее', all_text)
        self.assertNotIn('Богородский', all_text)  # original sample complaints were removed
        self.assertNotIn('99999', all_text)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_table:
                    self.assertLessEqual(shape.top + shape.height, prs.slide_height)

    def test_long_narratives_paginate_without_overflow(self):
        rows = presentation.validate_rows([record('ii', str(i), **{'Аннотация/краткое содержание': 'Длинное описание проблемы ' * 100, 'Ответ': 'Подробный ответ ведомства ' * 100}) for i in range(6)])
        prs = Presentation(io.BytesIO(presentation.build(rows, date(2026, 1, 1), date(2026, 1, 2))))
        count = 0
        for slide in list(prs.slides)[2:-1]:
            shape = presentation.named(slide, 'collective_rows')
            count += len(shape.table.rows) - 1
            self.assertLessEqual(shape.top + shape.height, prs.slide_height)
        self.assertEqual(count, 6)


if __name__ == '__main__':
    unittest.main()
