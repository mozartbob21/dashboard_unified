"""Isolated tests: no production startup, credentials or source calls."""
import io,json,sqlite3,time
from contextlib import contextmanager
import httpx
import pytest
from pathlib import Path
from fastapi.staticfiles import StaticFiles
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook
from services.omsu_rating import calculate as calc,source,store


def issue(omsu='ДУБНА',fields=None,name='РСО'):
    return {'id':1,'assigned_to':{'name':omsu},'subject':name,'custom_fields':[{'id':k,'value':v} for k,v in {17:1,**(fields or {})}.items()]}

def test_k4_categories_rounding_achievement_and_separate_zato():
    data=[issue('ДУБНА',{165:1000,169:100000}),issue('ИСТРА',{165:1000,169:200000}),issue('ВЛАСИХА',{165:1,169:1000}),issue('ФРЯЗИНО',{165:5000,169:300000})]
    rows={r['omsu']:r for r in calc.rating(4,0,data)}
    assert rows['ДУБНА']['place']==rows['ФРЯЗИНО']['place']==1
    assert rows['ИСТРА']['value']==50 and rows['ИСТРА']['place']==3
    assert rows['ВЛАСИХА']['place']==1 and rows['ВЛАСИХА']['group']=='zato'
    assert rows['ФРЯЗИНО']['details'][0]['plan']==1

def test_k4_second_rso_mean_and_absent_values_not_zero():
    row=calc.rating(4,1,[issue(fields={166:1000,169:100000,210:0,213:100000,208:'Вторая'})])[0]
    assert row['value']==25 and len(row['details'])==2
    assert calc.rating(4,0,[issue(fields={169:100000})])==[]
    assert calc.rating(4,0,[issue(fields={165:0,169:100000})])[0]['value']==0

def test_k5_zero_flags_ties_precise_values_and_missing():
    data=[issue('ДУБНА',{159:'6,6667'}),issue('ИСТРА',{159:'6,67'}),issue('ХИМКИ',{159:6.67}),issue('КЛИН',{159:9,259:'337'}),issue('ВЛАСИХА',{159:9,259:'337'}),issue('ФРЯЗИНО',{159:None})]
    rows={r['omsu']:r for r in calc.rating(5,1,data)}
    assert rows['ИСТРА']['place']==rows['ХИМКИ']['place']==1
    assert rows['ДУБНА']['place']==3
    assert rows['КЛИН']['value']==0 and rows['КЛИН']['place']==51
    assert rows['ВЛАСИХА']['place']==5
    assert 'ФРЯЗИНО' not in rows

def test_k6_conventional_place_and_exact_extra_totals():
    rows=calc.rating(6,2,[issue('ДУБНА',{201:20,196:100,197:10,198:5,199:3,200:2}),issue('ИСТРА',{201:2})])
    assert all(r['place']==1 for r in rows)
    row=next(r for r in rows if r['omsu']=='ДУБНА')
    assert row['extra']['total']==100 and row['extra']['without']==20
    assert row['extra']['parts']==[10,5,3,2]

def test_eligibility_requires_omsu_numeric17_and_keeps_finite_zero():
    assert not calc.rating(5,0,[issue('',{158:4}),issue(fields={17:None,158:4})])
    assert calc.num('0')==0 and calc.num(None) is None and calc.num('NaN') is None and calc.num(True) is None

def test_k7_provisional_rules_missing_and_duplicates():
    good=['']*22;good[2]='МУП «БКС»';good[3]='БАЛАШИХА'
    for item in calc.K7_ITEMS[:17]:good[item['col']]=100 if item['pct'] else 'выполнено'
    rows,_=calc.rating7([good,['','','МУП «БКС»','БАЛАШИХА']]);r=next(r for r in rows if r['omsu']=='БАЛАШИХА')
    assert r['value']==100 and r['details'][0]['duplicates']==2
    assert next(r for r in rows if r['omsu']=='ДУБНА')['value']==0
    assert not calc.k7cell('100%',{'n':5,'pct':True},True)['ok'] # archive conservative text rule
    assert not calc.k7cell('в работе',{'n':1,'pct':False},True)['ok']
    assert calc.k7cell(0,{'n':15,'pct':False},False)['ok']

def test_default_period_uses_coverage_and_k7_only_current9months():
    sources={'5':{'data':[issue('ДУБНА',{158:1,159:2,160:3,161:4}),issue('ИСТРА',{158:1,159:2,160:3}),issue('ХИМКИ',{158:1,159:2,160:3})]},'7':{'data':[]}}
    built=calc.build(sources);assert built['indicators']['5']['default_period']==2
    assert built['indicators']['7']['default_period']==2
    assert built['indicators']['7']['periods'][0]==[]

def test_redmine_fixed_target_header_pagination_and_privatefield_minimization():
    calls=[]
    def handler(request):
        calls.append(request);offset=int(request.url.params['offset'])
        row=issue(fields={158:1,999:'private-not-needed'});row['id']=offset+1;row['description']='not-used'
        return httpx.Response(200,json={'issues':[row],'total_count':2})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:values=source.fetch_redmine(52,'synthetic-key',client)
    assert len(values)==2 and calls[1].url.params['offset']=='1'
    assert all(r.headers['X-Redmine-API-Key']=='synthetic-key' for r in calls)
    assert all(str(r.url).startswith(source.REDMINE_ORIGIN+'/issues.json') for r in calls)
    assert 'description' not in values[0] and 999 not in [f['id'] for f in values[0]['custom_fields']]

@pytest.mark.parametrize('response',[httpx.Response(401,text='secret-upstream'),httpx.Response(500,text='secret-upstream'),httpx.Response(302,headers={'Location':'https://evil.test/secret'}),httpx.Response(200,json={'issues':[],'total_count':2})])
def test_upstream_errors_and_redirects_are_safe(response):
    with httpx.Client(transport=httpx.MockTransport(lambda _:response)) as client:
        with pytest.raises(source.SourceError) as exc:source.fetch_redmine(53,'secret-key',client)
    assert 'secret' not in str(exc.value) and 'evil.test' not in str(exc.value)

def test_sheet_parsing_and_no_token_on_google_requests():
    book=Workbook();sheet=book.active;sheet.title='Исполнение дорожной карты';sheet.append(['h']);sheet.append(['','','МУП «БКС»','БАЛАШИХА',1]);out=io.BytesIO();book.save(out)
    def handler(request):
        assert 'X-Redmine-API-Key' not in request.headers
        return httpx.Response(200,content=out.getvalue())
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:rows=source.fetch_sheet(client)
    assert rows[0][2]=='МУП «БКС»'

@pytest.fixture
def persisted(tmp_path,monkeypatch):
    monkeypatch.setattr(store,'DATA_DIR',tmp_path/'rating')
    from services.auth import source_api_keys
    monkeypatch.setattr(source_api_keys,'get_source_api_key',lambda _: 'synthetic-key')
    monkeypatch.setattr(source,'fetch_redmine',lambda q,k:[issue(fields={158:5,165:100,169:10000,189:10})])
    monkeypatch.setattr(source,'fetch_sheet',lambda:[])
    return tmp_path

def test_refresh_lease_lastgood_history_and_partial_failure(persisted,monkeypatch):
    token=store.claim();assert token and store.claim() is None
    store.refresh(token);saved=store.read();assert saved['refresh']['status']=='success' and len(saved['history'])==1
    def fail(*args):raise source.SourceError('Источник недоступен.')
    monkeypatch.setattr(source,'fetch_redmine',fail)
    store.refresh(store.claim());latest=store.read();assert latest['refresh']['status']=='partial'
    assert latest['indicators']['5']['periods']==saved['indicators']['5']['periods']
    assert latest['indicators']['5']['updated_at']==saved['indicators']['5']['updated_at']
    assert latest['indicators']['5']['error']=='Источник недоступен.'
    assert len(latest['history'])==2

def test_stale_worker_cannot_overwrite_newer_snapshot(persisted):
    old=store.claim()
    with store.connection() as conn:conn.execute('UPDATE refresh SET until=0 WHERE id=1')
    new=store.claim();assert new!=old
    store.refresh(old);assert store.read()['history']==[]
    store.refresh(new);assert len(store.read()['history'])==1

@pytest.fixture
def api(persisted,monkeypatch):
    from routers import omsu_rating
    from utils import db
    path=persisted/'accounts.sqlite'
    with sqlite3.connect(path) as conn:
        conn.execute('CREATE TABLE users(id INTEGER PRIMARY KEY,is_active INTEGER,modules TEXT)')
        conn.execute('CREATE TABLE account_archives(user_id INTEGER PRIMARY KEY)')
        conn.execute('INSERT INTO users VALUES(1,1,?)',(json.dumps(['omsu-rating']),))
    @contextmanager
    def connect():
        conn=sqlite3.connect(path);conn.row_factory=sqlite3.Row
        try:yield conn;conn.commit()
        finally:conn.close()
    monkeypatch.setattr(db,'get_db_connection',connect)
    app=FastAPI()
    app.mount('/static',StaticFiles(directory=Path(__file__).resolve().parents[1]/'static'),name='static')
    @app.middleware('http')
    async def actor(request,next):request.state.user={'id':1,'modules':['omsu-rating']};return await next(request)
    app.include_router(omsu_rating.router)
    with TestClient(app,base_url='https://neurona.test') as client:yield client,connect

def test_all_endpoints_enforce_fresh_grant_and_refresh_origin(api):
    client,connect=api
    assert client.get('/omsu-rating/api/data').status_code==200
    assert client.post('/omsu-rating/api/refresh',headers={'Origin':'https://evil.test'}).status_code==403
    assert client.post('/omsu-rating/api/refresh',headers={'Origin':'https://neurona.test'}).status_code==202
    with connect() as conn:conn.execute("UPDATE users SET modules='[]'")
    for path in ('','/api/data','/api/status','/export.xlsx'):assert client.get('/omsu-rating'+path).status_code==403
    assert client.post('/omsu-rating/api/refresh',headers={'Origin':'https://neurona.test'}).status_code==403

def test_export_real_workbook_all_indicators_and_formula_safe(api,monkeypatch):
    client,_=api
    monkeypatch.setattr(source,'fetch_redmine',lambda q,k:[issue('=HYPERLINK("evil")',{158:5,165:100,169:10000,189:10},'=1+1')])
    store.refresh(store.claim())
    response=client.get('/omsu-rating/export.xlsx');assert response.status_code==200
    book=load_workbook(io.BytesIO(response.content));assert len(book.sheetnames)==4
    assert book['Показатель 5']['C5'].value.startswith('=HYPERLINK')
    assert book['Показатель 5']['C5'].data_type=='s'
    assert response.headers['cache-control']=='private, no-store'
    assert client.get('/omsu-rating/export.xlsx?period=9').status_code==400

def test_snapshot_export_is_portable_escaped_and_has_no_credentials(api,monkeypatch):
    client,connect=api
    monkeypatch.setattr(source,'fetch_redmine',lambda q,k:[issue('<img src=x onerror=alert(1)>',{158:5,165:100,169:10000,189:10})])
    store.refresh(store.claim());response=client.get('/omsu-rating/export.html')
    assert response.status_code==200
    from test.test_omsu_rating_exports import SnapshotParser
    parsed=SnapshotParser();parsed.feed(response.text)
    assert len(parsed.scripts)==2 and not parsed.external
    assert '<img src=x onerror=alert(1)>' not in response.text
    assert '\\u003cimg' in response.text and 'synthetic-key' not in response.text
    assert "default-src 'none'" in parsed.policy and "connect-src 'none'" in parsed.policy
    with connect() as conn:conn.execute('UPDATE users SET is_active=0')
    assert client.get('/omsu-rating/export.html').status_code==403

def test_concurrent_refresh_requests_claim_exactly_one_lease(persisted):
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=6) as pool:results=list(pool.map(lambda _:store.claim(),range(12)))
    assert sum(token is not None for token in results)==1


def test_native_page_uses_platform_shell_local_assets_and_no_secret(api):
    client,_=api
    response=client.get('/omsu-rating')
    assert response.status_code==200
    assert 'Рейтинг муниципалитетов' in response.text
    assert '/static/omsu-rating.js' in response.text and '/static/module-shell.js' in response.text
    assert 'Снимок HTML' in response.text and 'Скачать Excel' in response.text
    assert 'synthetic-key' not in response.text and 'cdn.jsdelivr' not in response.text


@pytest.mark.parametrize('change',['duplicate','total','missing_id'])
def test_redmine_inconsistent_pages_never_double_count(change):
    calls=0
    def handler(request):
        nonlocal calls
        calls+=1;row=issue(fields={158:5});row['id']=1 if calls==1 or change=='duplicate' else 2
        if calls==2 and change=='missing_id':row.pop('id')
        return httpx.Response(200,json={'issues':[row],'total_count':3 if calls==2 and change=='total' else 2})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        with pytest.raises(source.SourceError):source.fetch_redmine(52,'synthetic-key',client)


def test_archived_active_account_and_malformed_modules_fail_closed(api):
    client,connect=api
    with connect() as conn:conn.execute('INSERT INTO account_archives VALUES(1)')
    assert client.get('/omsu-rating/api/data').status_code==403
    with connect() as conn:
        conn.execute('DELETE FROM account_archives');conn.execute('UPDATE users SET modules=?',(json.dumps('omsu-rating'),))
    assert client.get('/omsu-rating/export.xlsx').status_code==403


def test_k7_zero_raw_value_is_not_erased_in_detail():
    row=['']*22;row[2]='МУП «БКС»';row[3]='БАЛАШИХА';row[19]=0;row[20]=None
    rows,_=calc.rating7([row]);item=next(r for r in rows if r['omsu']=='БАЛАШИХА')
    details=item['details'][0]['items']
    assert details[14]['raw']=='0' and details[14]['ok'] is True
    assert details[15]['raw']=='' and details[15]['ok'] is False
