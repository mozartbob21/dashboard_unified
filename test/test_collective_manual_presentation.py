"""Shared manual reports remain complete, editable and literal in PPTX exports."""
import io
import re
import unittest
import zipfile
from datetime import date

from pptx import Presentation

from services.collective import presentation
from services.collective.client import HEADERS
from services.collective.presentation_text import status_blocks


def appeal(**manual):
    row = {key: '' for key in HEADERS}
    row.update({'_theme': 'ii', 'ОМСУ': 'Тестовый округ', 'Статус': 'В работе',
                'Аннотация/краткое содержание': 'Просим проверить водоснабжение.',
                'Ответ': '<p>Исходный ответ <b>ведомства</b>.</p>', HEADERS[10]: '101'})
    row.update(manual)
    return row


def build(rows):
    return presentation.build(presentation.validate_rows(rows), date(2026, 1, 1), date(2026, 1, 2))


def tables(prs):
    return [shape.table for slide in list(prs.slides)[2:-1]
            for shape in slide.shapes if shape.name == 'collective_rows']


class CollectiveManualPresentationTests(unittest.TestCase):
    def test_missing_fields_and_none_remain_empty(self):
        for row in (appeal(), appeal(_assignee=None, _report=None)):
            validated = presentation.validate_rows([row])[0]
            self.assertEqual(validated['_assignee'], '')
            self.assertEqual(validated['_report'], '')
            self.assertEqual(status_blocks(validated), [
                ('В работе', True), ('Исходный ответ ведомства.', False)])

    def test_manual_fields_preserve_literal_text_and_normalize_line_endings(self):
        report = '<b>Проверить</b> &lt;узел&gt;\r\n<script>буквальный текст</script>\rКонец'
        validated = presentation.validate_rows([appeal(
            _assignee='  Специалист <А>  ', _report='  ' + report + '  ')])[0]
        self.assertEqual(validated['_assignee'], 'Специалист <А>')
        expected = report.replace('\r\n', '\n').replace('\r', '\n')
        self.assertEqual(validated['_report'], expected)
        self.assertEqual(status_blocks(validated), [
            ('В работе', True), ('Исходный ответ ведомства.', False),
            ('Ответственный: Специалист <А>', True), ('Доклад', True), (expected, False)])

    def test_exact_maximum_lengths_are_kept(self):
        validated = presentation.validate_rows([appeal(
            _assignee='А' * 120, _report='Б' * 20000)])[0]
        self.assertEqual(len(validated['_assignee']), 120)
        self.assertEqual(len(validated['_report']), 20000)

    def test_invalid_manual_values_and_oversize_are_rejected(self):
        for field, limit in (('_assignee', 120), ('_report', 20000)):
            for value in (True, 7, [], {}, 'А' * (limit + 1)):
                with self.subTest(field=field, value_type=type(value).__name__):
                    with self.assertRaisesRegex(ValueError, 'ответственный|доклад'):
                        presentation.validate_rows([appeal(**{field: value})])

    def test_invalid_xml_characters_are_rejected_before_building(self):
        for field in ('_assignee', '_report'):
            for char in ('\x00', '\x01', '\ud800', '\ufffe', '\uffff'):
                with self.subTest(field=field, character=repr(char)):
                    with self.assertRaisesRegex(ValueError, 'недопустимые символы'):
                        presentation.validate_rows([appeal(**{field: 'Текст' + char})])

    def test_literal_markup_is_safe_text_in_presentation_xml(self):
        report = '<script>текст</script> &lt;узел&gt; <b>доклад</b> 2 < 3'
        data = build([appeal(_assignee='Специалист <А>', _report=report)])
        prs = Presentation(io.BytesIO(data))
        text = '\n'.join(row.cells[3].text for table in tables(prs) for row in list(table.rows)[1:])
        flattened = re.sub(r'\s+', '', text)
        self.assertIn(re.sub(r'\s+', '', report), flattened)
        self.assertIn('Ответственный:Специалист<А>', flattened)
        self.assertIn('Исходныйответведомства.', flattened)
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            xml = b''.join(package.read(name) for name in package.namelist()
                           if re.fullmatch(r'ppt/slides/slide\d+\.xml', name))
        self.assertIn(b'&lt;script&gt;', xml)
        self.assertIn(b'&amp;lt;', xml)
        self.assertNotIn(b'<script>', xml)
        self.assertNotIn(b'<b>', xml)

    def test_long_report_paginates_without_losing_any_words_or_inflating_totals(self):
        tokens = ['ДОКЛАД%04d' % index for index in range(1700)] + ['КОНЕЦ_ДОКЛАДА']
        report = ' '.join(tokens)
        self.assertLessEqual(len(report), 20000)
        prs = Presentation(io.BytesIO(build([appeal(_assignee='Проверочный специалист', _report=report)])))
        detail_tables = tables(prs)
        self.assertGreater(len(detail_tables), 1)
        text = '\n'.join(row.cells[3].text for table in detail_tables for row in list(table.rows)[1:])
        for token in tokens:
            self.assertEqual(text.count(token), 1, token)
        totals = presentation.named(prs.slides[1], 'collective_totals').table
        self.assertEqual(totals.cell(2, 0).text, '1')
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_table:
                    self.assertLessEqual(shape.left + shape.width, prs.slide_width)
                    self.assertLessEqual(shape.top + shape.height, prs.slide_height)


if __name__ == '__main__':
    unittest.main()
