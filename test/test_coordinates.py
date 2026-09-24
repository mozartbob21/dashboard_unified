import io
import json
from concurrent.futures import ThreadPoolExecutor

import pytest
from openpyxl import Workbook, load_workbook
from services.tools import coordinates as c
from services.tools.workspace import ToolError


def xlsx(rows, headers=('ID', 'Широта', 'Долгота', 'ОМСУ')):
    book = Workbook(); sheet = book.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'
    out = io.BytesIO(); book.save(out)
    return out.getvalue()


def run(tmp_path, rows, *, root=None, radius=500, prefix='VZ', name='test.xlsx', headers=None, imported=None):
    root = root or tmp_path/'owner'; root.mkdir(exist_ok=True)
    job = tmp_path/'job'; job.mkdir(exist_ok=True)
    data = xlsx(rows, headers) if headers else xlsx(rows)
    result = c.group(name, data, radius, prefix, root, job, imported)
    return result, load_workbook(job/result['file'])


@pytest.mark.parametrize('a,b,expected', [
    (55.75, 37.6, (55.75, 37.6)), ('37,6', '55,75', (55.75, 37.6)),
    ('55°45\'00"', '37°36\'00"', (55.75, 37.6)),
    ('55.75,37.6', None, (55.75, 37.6)), ('55,75; 37,6', None, (55.75, 37.6)),
    ('55 45', '37 36', (55.75, 37.6)), ('55.75 37.6', None, (55.75, 37.6)),
])
def test_coordinate_formats_and_swapping(a, b, expected):
    assert c.parse_coords(a,b)[:2] == pytest.approx(expected)


@pytest.mark.parametrize('a,b', [('-55.75','37.6'), ('55.75,-37.6',None),
                               ('55.75; -37.6',None), (float('inf'),37.6), ('bad','value'), (51,30)])
def test_bad_coordinates_are_not_fabricated(a,b):
    assert c.parse_coords(a,b)[0] is None


def test_groups_are_not_transitive_and_bad_rows_are_preserved(tmp_path):
    result,book = run(tmp_path, [[1,55.750,37.60,'Тест'],[2,55.753,37.60,'Тест'],[3,55.756,37.60,'Тест'],[4,'?','?',None]])
    assert result['objects'] == 2 and result['rows'] == 4 and result['unmatched'] == 1
    assert book.sheetnames == ['Объекты','Записи','Без координат','Реестр']
    headers = [x.value for x in book['Объекты'][1]]
    spread = headers.index('Разброс точек, м')
    assert all(row[spread] <= 500 for row in list(book['Объекты'].values)[1:])
    assert book['Без координат'].max_row == 2


def test_registry_codes_survive_new_files_and_are_account_private(tmp_path):
    first = tmp_path/'first'; first.mkdir()
    result,book = run(first,[[1,55.75,37.60,None]],prefix='',name='заявки тип VZ.xlsx')
    assert book['Записи']['A2'].value == 'VZ-0001'
    root = first/'owner'
    second = tmp_path/'second'; second.mkdir()
    _,book = run(second,[[2,55.751,37.60,None],[3,56.10,38.0,None]],root=root)
    assert {r[0] for r in list(book['Записи'].values)[1:]} == {'VZ-0001','VZ-0002'}
    other = tmp_path/'other'; other.mkdir()
    _,other_book = run(other,[[3,56.10,38.0,None]])
    assert other_book['Записи']['A2'].value == 'VZ-0001'
    assert len(json.loads((root/'coordinate_registry.json').read_text())) == 2


def test_import_export_registry_sheet_and_literal_excel_text(tmp_path):
    a=tmp_path/'a';a.mkdir()
    _,book = run(a, [[1,55.75,37.6,'=1+1']])
    uploaded=(a/'job/Объекты.xlsx').read_bytes()
    assert book['Записи']['E2'].value == '=1+1' and book['Записи']['E2'].data_type == 's'
    b=tmp_path/'b';b.mkdir()
    _,book = run(b, [[2,55.7501,37.6,None]], imported=uploaded)
    assert book['Записи']['A2'].value == 'VZ-0001'
    assert book['Реестр'].max_row == 2


def test_invalid_input_does_not_change_registry(tmp_path):
    run(tmp_path, [[1,55.75,37.6,'Тест']])
    registry=tmp_path/'owner/coordinate_registry.json'; original=registry.read_bytes()
    for radius in ('nan',0,5001):
        with pytest.raises(ToolError):
            c.group('bad.xlsx', xlsx([[1,55.75,37.6,'Тест']]), radius, 'VZ', registry.parent, tmp_path/'job')
    with pytest.raises(ToolError,match='Нет распознанных'):
        c.group('bad.xlsx', xlsx([[1,None,None,'Тест']]), 500, 'VZ', registry.parent, tmp_path/'job')
    assert registry.read_bytes() == original


def test_placeholder_marks_and_empty_municipality_do_not_crash(tmp_path):
    result,book=run(tmp_path, [[i,55.75,37.6,'Тест'] for i in range(6)] + [[7,56.1,38.0,None]])
    assert result['objects']==2
    assert any(row[1]=='условная точка округа' for row in list(book['Объекты'].values)[1:])


def test_same_account_parallel_jobs_allocate_distinct_codes(tmp_path):
    root=tmp_path/'owner';root.mkdir()
    jobs=[tmp_path/'a',tmp_path/'b']
    for job in jobs:job.mkdir()
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures=[pool.submit(c.group,'data.xlsx',xlsx([[i,55.7+i*.1,37.6,None]]),500,'VZ',root,job) for i,job in enumerate(jobs)]
        for f in futures:assert f.result()['objects']==1
    saved=json.loads((root/'coordinate_registry.json').read_text())
    assert {r['Код объекта'] for r in saved}=={'VZ-0001','VZ-0002'}
