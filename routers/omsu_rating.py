"""Native rating module; all data, jobs and exports require a current grant."""
import io
from html import escape
import json
from fastapi import APIRouter,BackgroundTasks,Depends,HTTPException,Request
from fastapi.responses import JSONResponse,Response
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from core.web import templates
from services.omsu_rating import store,calculate

PRIVATE={'Cache-Control':'private, no-store'}

def require_rating(request:Request):
    user=getattr(request.state,'user',None)
    if not user:
        from services.auth.security import get_user_from_token
        user=get_user_from_token(request.cookies.get('access_token'))
    if not user:raise HTTPException(401,'Войдите в Нейрону.')
    if user.get('kc_sub'):
        from core.roles import check_module_access
        allowed=check_module_access(user,'omsu-rating')
    else:
        from utils.db import get_db_connection
        if type(user.get('id')) is not int or user['id']<=0:raise HTTPException(401,'Войдите в Нейрону.')
        with get_db_connection() as conn:
            row=conn.execute('SELECT is_active,modules FROM users WHERE id=? AND NOT EXISTS (SELECT 1 FROM account_archives WHERE user_id=users.id)',(user['id'],)).fetchone()
        try:
            modules=json.loads(row['modules'] or '[]') if row else []
            allowed=bool(row and row['is_active'] and isinstance(modules,list) and 'omsu-rating' in modules)
        except (ValueError,TypeError):allowed=False
    if not allowed:raise HTTPException(403,'Нет доступа к рейтингу ОМСУ.')
    return user

router=APIRouter(prefix='/omsu-rating',dependencies=[Depends(require_rating)])

@router.get('')
def page(request:Request):return templates.TemplateResponse(request,'omsu_rating.html',{'request':request},headers=PRIVATE)

@router.get('/api/data')
def data():return JSONResponse(store.read(),headers=PRIVATE)

@router.get('/api/status')
def status():return JSONResponse(store.status(),headers=PRIVATE)

@router.post('/api/refresh')
def refresh(request:Request,background:BackgroundTasks):
    origin=request.headers.get('origin','')
    expected=f'{request.url.scheme}://{request.url.netloc}'
    if origin!=expected or request.headers.get('sec-fetch-site')=='cross-site':raise HTTPException(403,'Обновление доступно только со страницы Нейроны.')
    token=store.claim()
    if token is not None:background.add_task(store.refresh,token)
    return JSONResponse({'status':'running','already_running':token is None},status_code=202,headers=PRIVATE)

@router.get('/export.xlsx')
def export(ind:int|None=None,period:int|None=None):
    if ind is not None and ind not in (4,5,6,7):raise HTTPException(400,'Неизвестный показатель.')
    if period is not None and period not in range(4):raise HTTPException(400,'Неизвестный период.')
    saved=store.read();book=Workbook();book.remove(book.active)
    for key in (4,5,6,7):
        item=saved['indicators'][str(key)];p=period if ind==key and period is not None else item['default_period']
        sheet=book.create_sheet(f'Показатель {key}');sheet.append([item['name']]);sheet.append([item['note']]);sheet.append(['Период: '+(calculate.PERIODS[p] if p is not None else 'Нет данных')+' · Получено: '+(item['updated_at'] or 'Не получены')])
        sheet.append(['Место','ОМСУ','Группа',item['metric'],'Факт / план, %','ЛС без ИПУ / всего','Мероприятия / план','РСО / задачи','Примечание'])
        if p is not None:
            for row in item['periods'][p]:
                extra=row['extra']
                sheet.append([row['place'],row['omsu'],'ЗАТО' if row['group']=='zato' else 'ОМСУ',row['value'],
                    str(extra['fact'])+' / '+', '.join(map(str,extra['plan'])) if 'fact' in extra else '',
                    str(extra['without'])+' / '+str(extra['total']) if 'total' in extra else '',
                    str(extra['done'])+' / '+str(extra['plan']) if 'done' in extra else '',
                    '; '.join(x['name'] for x in row['details']),'Балл обнулён' if row['no_docs'] else ''])
        if p is None or not item['periods'][p]:sheet.append(['Пока нет данных'])
        # Source-supplied strings stay literal even when beginning with '='.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in sheet[4]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='515DCE')
        for col,width in [('A',12),('B',29),('C',18),('D',26),('E',24),('F',24),('G',24),('H',55),('I',25)]:sheet.column_dimensions[col].width=width
        for title_row in (1,2,3):sheet.merge_cells(start_row=title_row,start_column=1,end_row=title_row,end_column=9)
        sheet.row_dimensions[1].height=28;sheet.row_dimensions[2].height=42;sheet.row_dimensions[3].height=24
        sheet.freeze_panes='C5';sheet.auto_filter.ref=f'A4:I{sheet.max_row}'
    out=io.BytesIO();book.save(out)
    return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={**PRIVATE,'Content-Disposition':'attachment; filename="omsu-rating.xlsx"'})


@router.get('/export.html')
def snapshot():
    """Portable safe report of cached ratings; no tokens, scripts or external requests."""
    saved=store.read();parts=[]
    for key in (4,5,6,7):
        item=saved['indicators'][str(key)];period=item['default_period']
        parts.append('<section><h2>'+escape(item['name'])+'</h2><p>'+escape(item['note'])+'</p><p>Получено: '+escape(item['updated_at'] or 'Нет данных')+'</p>')
        if period is None:parts.append('<p>Пока нет данных.</p></section>');continue
        parts.append('<p>Период: '+escape(calculate.PERIODS[period])+'</p><table><thead><tr><th>Место</th><th>ОМСУ</th><th>Группа</th><th>'+escape(item['metric'])+'</th></tr></thead><tbody>')
        for row in item['periods'][period]:
            parts.append('<tr>'+''.join('<td>'+escape(str(value))+'</td>' for value in [row['place'],row['omsu'],'ЗАТО' if row['group']=='zato' else 'ОМСУ',round(row['value'],2)])+'</tr>')
        parts.append('</tbody></table></section>')
    html='''<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'"><title>Рейтинг ОМСУ — снимок</title><style>body{font:15px/1.6 system-ui,sans-serif;max-width:1100px;margin:36px auto;padding:0 20px;color:#1d293e}section{margin:28px 0;padding:24px;border:1px solid #dde3ed;border-radius:16px;overflow:auto}table{width:100%;border-collapse:collapse}td,th{text-align:left;padding:10px;border-bottom:1px solid #dde3ed}th{background:#f1f3fc}p{color:#5b677a}@media print{section{break-before:page;border:0}h1{margin-top:0}}</style><h1>Рейтинг ОМСУ · Нейрона ИИ</h1><p>Методика 2026 года. Сохранённый снимок; данные не обновляются автоматически.</p>'''+''.join(parts)+'</html>'
    return Response(html,media_type='text/html',headers={**PRIVATE,'Content-Disposition':'attachment; filename="omsu-rating-snapshot.html"','X-Content-Type-Options':'nosniff'})
