"""Editable collective-appeal slides using the supplied Ministry template.

Text is measured and continued on extra pages, never silently abbreviated.
Only local data and fonts are used. No executable code from the archive is run.
"""
import copy
import io
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from services.collective.client import HEADERS
from services.collective.presentation_text import clean, content_blocks, status_blocks
from services.collective import presentation_layout as layout

TEMPLATE = Path(__file__).parent / 'assets' / 'template.pptx'
THEMES = {'ii': 'Инженерная инфраструктура', 'kr': 'Капитальный ремонт МКД', 'pr': 'Прочее'}
MAX_ROWS = 2000
MAX_CELL = 100000
MAX_SLIDES = 500
DARK = '172033'


def signature(value):
    match = re.match(r'^\+?(\d+)', re.sub(r'\s', '', str(value or '')))
    if not match:
        return None
    digits = match[1].lstrip('0') or '0'
    if len(digits) > 16 or int(digits) > 2**53 - 1:
        raise ValueError('Слишком большое число подписей. Проверьте исходные данные.')
    return int(digits)


def validate_rows(rows):
    if not isinstance(rows, list) or not rows:
        raise ValueError('Нет обращений для презентации.')
    if len(rows) > MAX_ROWS:
        raise ValueError(f'В презентацию можно включить до {MAX_ROWS} обращений. Уточните фильтры.')
    result = []
    for row in rows:
        if not isinstance(row, dict) or row.get('_theme') not in THEMES:
            raise ValueError('Некорректная тема обращения.')
        item = {}
        for key in HEADERS:
            value = row.get(key, '')
            if value is None:
                value = ''
            if isinstance(value, (list, dict, bool)) or len(str(value)) > MAX_CELL:
                raise ValueError('Некорректные поля обращения.')
            item[key] = str(value)
        item['_theme'] = row['_theme']
        item['_sig'] = signature(item[HEADERS[10]])
        result.append(item)
    return result



def named(slide, name):
    return next(shape for shape in slide.shapes if shape.name == name)


def clone_slide(prs, source):
    slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    rels = {}
    for rid, rel in source.part.rels.items():
        if rel.reltype.endswith(('/slideLayout', '/notesSlide')):
            continue
        rels[rid] = slide.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,
                                       rel.reltype, is_external=rel.is_external)
    for shape in source.shapes:
        element = copy.deepcopy(shape.element)
        for child in element.iter():
            for attr in ('r:embed', 'r:id', 'r:link'):
                old = child.get(qn(attr))
                if old in rels:
                    child.set(qn(attr), rels[old])
        slide.shapes._spTree.insert_element_before(element, 'p:extLst')
    return slide



def font_style(paragraph, size=14, bold=False, color=DARK):
    paragraph.font.name = 'Arial'
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)


def text_cell(cell, text, *, bold=False, size=14, centered=False):
    cell.text = str(text)
    cell.margin_left = cell.margin_right = Inches(layout.CELL_MARGIN_X)
    cell.margin_top = cell.margin_bottom = Inches(layout.CELL_MARGIN_Y)
    cell.vertical_anchor = MSO_ANCHOR.TOP
    cell.text_frame.word_wrap = True
    cell.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
        font_style(paragraph, size, bold)
        paragraph.line_spacing = Pt(size * 1.2)


def shape_text(shape, value, size, *, bold=False, color=DARK):
    shape.text = str(value)
    shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    for p in shape.text_frame.paragraphs:
        font_style(p, size, bold, color)


def rich_cell(cell, lines, centered=False):
    text_cell(cell, '', centered=centered)
    tf = cell.text_frame
    for i, (line, bold, gap) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
        font_style(p, layout.FONT_SIZE, bold)
        p.line_spacing = Pt(layout.LINE_HEIGHT)
        p.space_after = Pt(gap)


def details(row):
    return ('\n'.join(text for text, _ in content_blocks(row)),
            '\n'.join(text for text, _ in status_blocks(row)))


def row_cells(row):
    fields = [[(clean(row.get('ОМСУ')) or 'Не указан', True)],
              [(str(row['_sig']) if row['_sig'] is not None else 'Не указано', False)],
              content_blocks(row), status_blocks(row)]
    return [layout.lines_for(blocks, i) for i, blocks in enumerate(fields)]


def row_height(row):
    return layout.row_height(row_cells(row))


def fill_summary(slide, rows):
    totals = named(slide, 'collective_totals').table
    labels = ['Поступило', 'Более 100 подписей', 'До 100 включительно', 'Повторных']
    known = [r for r in rows if r['_sig'] is not None]
    repeats = [r for r in rows if clean(r['Повтор']).lower() == 'да']
    values = [len(rows), sum(r['_sig'] > 100 for r in known),
              sum(r['_sig'] <= 100 for r in known), len(repeats)]
    text_cell(totals.cell(0, 0), 'Всего обращений', bold=True, size=18)
    for i, (label, value) in enumerate(zip(labels, values)):
        totals.columns[i].width = Inches(12.34 / 4)
        text_cell(totals.cell(1, i), label, bold=True, centered=True)
        text_cell(totals.cell(2, i), value, size=26, centered=True)
    repeat_places = sorted({clean(r['ОМСУ']) for r in repeats if clean(r['ОМСУ'])})
    if repeat_places:
        # Keep the full roster in editable notes; show a bounded preview in the cell.
        roster = ', '.join(repeat_places)
        roster_lines = layout.wrap_text(roster, 12.34 / 4, size=11)
        preview = '\n'.join(roster_lines) if len(roster_lines) <= 2 else (
            roster_lines[0].rstrip('., ') + '…\nПеречень в примечаниях')
        p = totals.cell(2, 3).text_frame.add_paragraph()
        p.text = preview
        p.alignment = PP_ALIGN.CENTER
        font_style(p, 11)
        p.line_spacing = Pt(13)
        slide.notes_slide.notes_text_frame.text = 'Повторные обращения. ОМСУ: ' + roster
    for row, height in zip(totals.rows, (.48, .55, 1.0)):
        row.height = Inches(height)
    unknown = len(rows) - len(known)
    if unknown:
        note = slide.shapes.add_textbox(Inches(.16), Inches(2.95), Inches(12.3), Inches(.28))
        note.name = 'collective_unknown_signatures'
        note.text_frame.margin_top = note.text_frame.margin_bottom = 0
        shape_text(note, f'Число подписей не указано: {unknown}. Эти обращения не включены в группы по числу подписей.', 12)
    theme_shape = named(slide, 'collective_themes')
    theme_shape.top = Inches(3.38)
    table = theme_shape.table
    while len(table.rows) < 5:
        table._tbl.append(copy.deepcopy(table._tbl.findall(qn('a:tr'))[-1]))
    text_cell(table.cell(0, 0), 'Учитываемые обращения по темам', bold=True, size=18)
    for i, label in enumerate(['Категория', 'Общее количество', 'Более 100 подписей']):
        text_cell(table.cell(1, i), label, bold=True, centered=i > 0)
    active = [r for r in rows if 'не учитывается' not in clean(r['Статус']).lower()]
    for i, (key, title) in enumerate(THEMES.items(), 2):
        group = [r for r in active if r['_theme'] == key]
        for j, value in enumerate([title, len(group), sum(r['_sig'] is not None and r['_sig'] > 100 for r in group)]):
            text_cell(table.cell(i, j), value, bold=j == 0, size=15, centered=j > 0)
    for row, height in zip(table.rows, (.48, .6, .72, .72, .72)):
        row.height = Inches(height)


def paginate(rows):
    # Preserve the source generator's priority: signatures across the full export,
    # with no forced half-empty page at each change of topic.
    ordered = sorted(rows, key=lambda r: r['_sig'] if r['_sig'] is not None else -1, reverse=True)
    themes = {r['_theme'] for r in rows}
    label = THEMES[next(iter(themes))] if len(themes) == 1 else 'Коллективные обращения'
    pages, chunk, used = [], [], 0
    for row in ordered:
        for cells in layout.fragments(row_cells(row)):
            height = layout.row_height(cells)
            if chunk and (used + height > layout.BODY_HEIGHT or len(chunk) >= layout.MAX_ROWS_PER_SLIDE):
                pages.append((label, chunk))
                chunk, used = [], 0
            chunk.append(cells)
            used += height
            if len(pages) + 4 > MAX_SLIDES:
                raise ValueError(f'Презентация превышает {MAX_SLIDES} слайдов. Уточните фильтры или разделите период.')
    if chunk:
        pages.append((label, chunk))
    if len(pages) + 3 > MAX_SLIDES:
        raise ValueError(f'Презентация превышает {MAX_SLIDES} слайдов. Уточните фильтры или разделите период.')
    return pages


def build(rows, start, end):
    pages = paginate(rows)
    prs = Presentation(TEMPLATE)
    cover, summary, sample, thanks = list(prs.slides)
    period = f'{start:%d.%m.%Y} – {end:%d.%m.%Y}'
    shape_text(named(cover, 'collective_period'), 'Свод ' + period, 32, bold=True, color='FFFFFF')
    shape_text(named(summary, 'collective_title'), 'Коллективные обращения ' + period, 24, bold=True, color='FFFFFF')
    fill_summary(summary, rows)
    generated = []
    for label, chunk in pages:
        slide = clone_slide(prs, sample)
        generated.append(slide)
        title = slide.shapes.add_textbox(Inches(.16), Inches(.17), Inches(12.3), Inches(.55))
        title.name = 'collective_detail_title'
        shape_text(title, label + ' · ' + period, 22, bold=True)
        table_shape = named(slide, 'collective_rows')
        table_shape.top = Inches(.85)
        table = table_shape.table
        for column, width in zip(table.columns, layout.COLUMN_WIDTHS):
            column.width = Inches(width)
        body = copy.deepcopy(table._tbl.findall(qn('a:tr'))[1])
        for tr in list(table._tbl.findall(qn('a:tr')))[1:]:
            table._tbl.remove(tr)
        table.rows[0].height = Inches(layout.HEADER_HEIGHT)
        for cell in table.rows[0].cells:
            text_cell(cell, cell.text, bold=True, size=12)
        for cells in chunk:
            table._tbl.append(copy.deepcopy(body))
            out = table.rows[len(table.rows) - 1]
            out.height = Inches(layout.row_height(cells))
            for i, lines in enumerate(cells):
                rich_cell(out.cells[i], lines, centered=i == 1)
        table_shape.height = sum(row.height for row in table.rows)
    ids = {id(slide): element for slide, element in zip(prs.slides, list(prs.slides._sldIdLst))}
    sample_id = ids[id(sample)]
    prs.part.drop_rel(sample_id.rId)
    prs.slides._sldIdLst.remove(sample_id)
    order = [cover, summary, *generated, thanks]
    for element in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(element)
    for slide in order:
        prs.slides._sldIdLst.append(ids[id(slide)])
    for i, slide in enumerate(order, 1):
        markers = [shape for shape in slide.shapes if shape.name == 'collective_page']
        for duplicate in markers[1:]:
            duplicate._element.getparent().remove(duplicate._element)
        if markers:
            marker = markers[0]
            marker.left, marker.top = Inches(12.75), Inches(7.04)
            marker.width, marker.height = Inches(.48), Inches(.35)
            marker.text_frame.margin_left = marker.text_frame.margin_right = 0
            marker.text_frame.margin_top = marker.text_frame.margin_bottom = 0
            shape_text(marker, str(i), 14, bold=True, color='FFFFFF')
            marker.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
