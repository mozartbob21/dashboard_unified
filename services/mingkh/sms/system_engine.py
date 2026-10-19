#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Обновление реестра системных улиц (МинЖКХ МО).

Системная улица — улица, по которой за последние 30 дней поступило
не менее 3 жалоб минимум в 2 разных календарных дня.

Reviewed server module imported by sys_data; no executable content is loaded
from uploaded spreadsheets or HTML. The complaint year follows ref_date.
"""

import re
from datetime import timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Параметры критерия системности
MIN_COMPLAINTS = 3      # минимум жалоб за окно
MIN_DAYS = 2            # минимум разных календарных дней
WINDOW_DAYS = 30        # длина окна (включительно, считая дату отсчёта)

# ---------------------------------------------------------------------------
# Нормализация адресов и ключей

def norm(s):
    """Нижний регистр, ё→е, схлопывание пробелов, отсечение хвостовых запятых."""
    if pd.isna(s):
        return ''
    s = re.sub(r'\s+', ' ', str(s).lower().replace('ё', 'е')).strip()
    return s.strip(',').strip()

HOUSE_RE = re.compile(
    r'^(\d+[а-яa-z]?([/-]\d+[а-яa-z]?)?'
    r'((к|корп|корпус|с|стр|строение|лит|литера)\.?\s*\d*[а-яa-z]?)*'
    r'|(к|корп|корпус)\.?\s*\d+[а-яa-z]?'
    r'|владение\s*\d+\w*|вл\.?\s*\d+\w*)$',
    re.I,
)

def split_addr(full):
    """Полный адрес -> (ключ_улицы, отображаемая_улица, дом)."""
    if pd.isna(full):
        return None, None, None
    s = re.sub(r'^россия,\s*', '', str(full).strip(), flags=re.I)
    parts = [p.strip() for p in s.split(',') if p.strip()]
    house_parts = []
    while parts and HOUSE_RE.match(parts[-1].replace(' ', '')):
        house_parts.insert(0, parts.pop())
    house = ' '.join(house_parts) if house_parts else None
    if not parts:
        return None, None, None
    disp = ', '.join(parts)
    return norm(disp), disp, house

def house_key(h):
    return norm(h).replace(' ', '').replace('.', '') if (h and pd.notna(h)) else ''

# ---------------------------------------------------------------------------
# Сопоставление улиц системки с ключами жалоб (с разведением по ОМСУ)

OKRUG = re.compile(r'(городской округ|муниципальный округ|муниципальный район|городское поселение)')
ADMIN = re.compile(r'(ский|ой)$')

def _segs(key):
    return [p.strip() for p in key.split(',') if p.strip() and p.strip() != 'московская область']

def _core(key):
    return [p for p in _segs(key) if not OKRUG.search(p)]

def _tail(key, k=2):
    cs = [p for p in _core(key) if not (ADMIN.search(p) and ' ' not in p)]
    return tuple(cs[-k:]) if len(cs) >= k else tuple(cs)

def build_mapping(street_keys_with_omsu, complaints_df):
    """street_keys_with_omsu: список (ключ_улицы, омсу_норм).
    Возвращает dict: ключ_системки -> set(ключей_жалоб)."""
    fk_oms = complaints_df.groupby('ул_ключ')['омсу_н'].agg(lambda s: set(s)).to_dict()
    full_keys = set(complaints_df['ул_ключ'].dropna())

    idx2, idx1 = {}, {}
    for fk in full_keys:
        idx2.setdefault(_tail(fk, 2), set()).add(fk)
        t1 = _tail(fk, 1)
        if t1:
            idx1.setdefault(t1[0], set()).add(fk)

    mapping, unmatched = {}, []
    for k, oms in street_keys_with_omsu:
        if k in full_keys and oms in fk_oms.get(k, set()):
            mapping.setdefault(k, set()).add(k)
            continue
        cand = [fk for fk in full_keys if fk.endswith(', ' + k)]
        if not cand:
            core = set(_core(k))
            cand = [fk for fk in full_keys if core and core <= set(_core(fk))]
        if not cand:
            t2 = _tail(k, 2)
            if len(t2) == 2:
                cand = list(idx2.get(t2, []))
        if not cand:
            t1 = _tail(k, 1)
            if t1:
                cand = list(idx1.get(t1[0], []))
        # фильтр по ОМСУ — не даём слипнуться одноимённым улицам из разных округов
        cand = [fk for fk in cand if oms in fk_oms.get(fk, set())]
        if cand:
            mapping.setdefault(k, set()).update(cand)
        else:
            unmatched.append((oms, k))
    return mapping, unmatched

# ---------------------------------------------------------------------------
# Определение "настоящей" улицы (для отсечения записей уровня города)

STREET_MARKERS = re.compile(
    r'(улиц|проезд|переул|шоссе|проспект|бульвар|тупик|набережн|площад|аллея|'
    r'лини|тракт|квартал|микрорайон|деревн|село|поселок|пгт|снт|жк |тсн|днп|'
    r'дп |кп |территор|хутор|слобод|станци|городок|вал|спуск|съезд|просек|дорог)'
)

def is_real_street(disp):
    seg = disp.split(',')[-1].strip().lower().replace('ё', 'е')
    return bool(STREET_MARKERS.search(seg)) or bool(re.match(r'^\d', seg))

# ---------------------------------------------------------------------------
# Загрузка входных файлов

def load_complaints(path, ref_date):
    """Читает файл жалоб. Поддерживает файл с заголовками и без."""
    with pd.ExcelFile(path) as xls:
        sheet = 'Sheet1' if 'Sheet1' in xls.sheet_names else xls.sheet_names[0]
    probe = pd.read_excel(path, sheet_name=sheet, nrows=1, header=None)
    if probe.empty or len(probe.columns) < 4:
        raise ValueError('В выгрузке жалоб отсутствуют необходимые колонки.')
    has_header = any('дата' in str(v).lower() for v in probe.iloc[0])

    if has_header:
        raw = pd.read_excel(path, sheet_name=sheet)
        col = {str(c).lower(): c for c in raw.columns}
        required = ('дата', 'тип сообщения', 'адрес', 'омсу')
        if any(not any(token in name for name in col) for token in required):
            raise ValueError('В выгрузке жалоб отсутствуют дата, тип, адрес или ОМСУ. Обновление отменено.')
        def g(*names):
            for nm in names:
                for lc, orig in col.items():
                    if nm in lc:
                        return raw[orig]
            return pd.Series([None] * len(raw))
        df = pd.DataFrame({
            'ецур': g('номер ецур'),
            'источник': g('номер в источнике'),
            'дата': pd.to_datetime(g('дата (первого', 'дата взятия', 'дата'), errors='coerce', format='mixed'),
            'тип': g('тип сообщения'),
            'Адрес': g('адрес'),
            'Дом': g('дом'),
            'омсу_src': g('омсу'),
        })
    else:
        raw = pd.read_excel(path, sheet_name=sheet, header=None)
        if len(raw.columns) < 23:
            raise ValueError('В выгрузке жалоб без заголовков недостаточно колонок.')
        df = pd.DataFrame({
            'ецур': raw[1], 'источник': raw[2],
            'дата': pd.to_datetime(raw[3], errors='coerce', format='mixed'),
            'тип': raw[14], 'Адрес': raw[18], 'Дом': raw[22],
            'омсу_src': raw[11],
        })

    if len(df) and df['дата'].isna().any():
        raise ValueError('Не удалось распознать даты в выгрузке жалоб. Обновление отменено.')
    # Год отчёта плюс полное окно30дней: в январе оно включает декабрь.
    # Тип может быть числом (0/1) или текстом ("0"/"Проблемы"/"1"/"Предложения").
    def _is_problem(v):
        s = str(v).strip().lower()
        return s in ("0", "0.0", "проблемы", "проблема")
    ref_date = pd.Timestamp(ref_date).normalize()
    first_day = min(ref_date.replace(month=1, day=1), ref_date - timedelta(days=WINDOW_DAYS - 1))
    df = df[(df['дата'] >= first_day)
            & (df['дата'] < ref_date + timedelta(days=1))
            & df['тип'].apply(_is_problem)].copy()
    df['дата'] = df['дата'].dt.normalize()
    df['отчётный_год'] = df['дата'].dt.year == ref_date.year

    res = df['Адрес'].apply(split_addr)
    df['ул_ключ'] = res.apply(lambda x: x[0])
    df['ул_дисп'] = res.apply(lambda x: x[1])
    df['дом'] = res.apply(lambda x: x[2])
    no_h = df['дом'].isna() & df['Дом'].notna()
    df.loc[no_h, 'дом'] = df.loc[no_h, 'Дом'].astype(str)

    def num(r):
        v = r['источник']
        if pd.notna(v) and str(v).strip() not in ('', 'nan'):
            return re.sub(r'\.0$', '', str(v).strip())
        v = r['ецур']
        return re.sub(r'\.0$', '', str(v).strip()) if pd.notna(v) else None
    df['номер'] = df.apply(num, axis=1) if len(df) else pd.Series(dtype='object')
    df['омсу_н'] = df['омсу_src'].apply(norm)
    df['сег'] = df['ул_дисп'].apply(lambda d: d.split(',')[-1].strip() if pd.notna(d) else '')
    df['а_ключ'] = df['сег'].apply(norm) + '||' + df['дом'].apply(house_key)
    df['а_дисп'] = df.apply(
        lambda r: r['сег'] + (', ' + str(r['дом']) if pd.notna(r['дом']) and r['дом'] else ''), axis=1) if len(df) else pd.Series(dtype='object')
    return df

def load_systemka(path):
    with pd.ExcelFile(path) as xls:
        str_sheet = 'Sheet1' if 'Sheet1' in xls.sheet_names else xls.sheet_names[0]
        adr_sheet = 'Лист1' if 'Лист1' in xls.sheet_names else None
    s_str = pd.read_excel(path, sheet_name=str_sheet)
    # приводим статусную колонку "Статус на ДД.ММ.ГГГГ" к "Статус",
    # чтобы сравнение с прошлой неделей работало, даже если на вход подан
    # готовый результат предыдущего расчёта
    if 'Статус' not in s_str.columns:
        for c in s_str.columns:
            if str(c).startswith('Статус'):
                s_str = s_str.rename(columns={c: 'Статус'})
                break
    s_str['ул_ключ'] = s_str['Улица'].apply(norm)
    if adr_sheet:
        s_adr = pd.read_excel(path, sheet_name=adr_sheet)
    else:
        # если листа адресов нет — строим из листа улиц (одна улица = одна строка-агрегат)
        s_adr = s_str.rename(columns={}).copy()
        s_adr['Адрес'] = s_adr['Улица']
    # проверка целостности: на листе адресов должно быть заметно больше строк,
    # чем улиц на Sheet1 (у каждой улицы обычно несколько адресов).
    # Если адресов подозрительно мало — файл, вероятно, обрезан.
    n_streets = int(s_str['Улица'].nunique())
    n_addr_streets = int(pd.Series(s_adr['Улица']).nunique()) if 'Улица' in s_adr.columns else 0
    load_systemka.warning = ''
    if n_streets >= 50 and n_addr_streets < n_streets * 0.5:
        load_systemka.warning = (
            f'Похоже, файл системки неполный: на листе улиц {n_streets} улиц, '
            f'а на листе адресов только {n_addr_streets}. '
            f'Лист адресов (Лист1), вероятно, обрезан — возьмите полный результат прошлой недели.'
        )
    s_adr['ул_ключ'] = s_adr['Улица'].apply(norm)
    s_adr['дом_к'] = s_adr['Адрес'].apply(
        lambda a: house_key(str(a).rsplit(',', 1)[1])
        if pd.notna(a) and ',' in str(a) and re.match(r'^\s*\d', str(a).rsplit(',', 1)[1]) else '')
    s_adr['сег_к'] = s_adr.apply(
        lambda r: norm(str(r['Адрес']).rsplit(',', 1)[0]) if r['дом_к'] else norm(r['Адрес']), axis=1)
    s_adr['а_ключ'] = s_adr['сег_к'] + '||' + s_adr['дом_к']
    return s_str, s_adr

# ---------------------------------------------------------------------------
# Основной расчёт

def aggregate(group):
    g = group.sort_values('дата')
    nums = list(dict.fromkeys(g['номер'].dropna()))
    return pd.Series({
        'кол': len(g),
        'кол_год': int(g['отчётный_год'].sum()),
        'номера': ', '.join(map(str, nums)),
        'посл': g['дата'].max(),
        'дисп': g['а_дисп'].iloc[-1],
    })

def process(systemka_path, complaints_path, ref_date):
    ref_date = pd.Timestamp(ref_date).normalize()
    win_start = ref_date - timedelta(days=WINDOW_DAYS - 1)
    new_period = f"{win_start.strftime('%d.%m')}.-{ref_date.strftime('%d.%m.%Y')}"

    def status(d):
        return 'Жалобы продолжаются' if (pd.notna(d) and pd.Timestamp(win_start) <= pd.Timestamp(d) <= ref_date) \
            else 'Нет жалоб последние 30 дней'

    df = load_complaints(complaints_path, ref_date)
    s_str, s_adr = load_systemka(systemka_path)

    # сопоставление улиц
    street_keys = list(zip(s_str['ул_ключ'], s_str['ОМСУ'].apply(norm)))
    mapping, unmatched = build_mapping(street_keys, df)

    fk2sk = {}
    for skk, fks in mapping.items():
        for fk in fks:
            fk2sk.setdefault(fk, set()).add(skk)
    registry_omsu = dict(zip(s_str['ул_ключ'], s_str['ОМСУ'].apply(norm)))
    df['sk'] = df.apply(
        lambda row: {key for key in fk2sk.get(row['ул_ключ'], set())
                     if registry_omsu.get(key) == row['омсу_н']}, axis=1
    ) if len(df) else pd.Series(dtype='object')

    # метаданные улиц из листа адресов
    meta = {}
    for skk, grp in s_adr.groupby('ул_ключ', sort=False):
        b = grp.iloc[0]
        meta[skk] = dict(
            Улица=b['Улица'], ОМСУ=b['ОМСУ'],
            задачи=b.get('Номера задач #', ''),
            ш=b.get('Широта', b.get('широта', '')),
            д=b.get('Долгота', b.get('долгота', '')),
        )

    # ----- обновление адресов существующих улиц -----
    rows = []
    for skk in s_adr['ул_ключ'].unique():
        m = meta[skk]
        sub = df[df['sk'].apply(lambda s: bool(s) and skk in s).astype(bool)]
        ag = sub.groupby('а_ключ').apply(aggregate, include_groups=False) if len(sub) else pd.DataFrame()
        grp = s_adr[s_adr['ул_ключ'] == skk]
        used, out = set(), []
        for _, r in grp.iterrows():
            ak = r['а_ключ']
            if len(ag) and ak in ag.index:
                a = ag.loc[ak]; used.add(ak)
                out.append(dict(Улица=m['Улица'], Адрес=r['Адрес'], ОМСУ=m['ОМСУ'], кол=int(a['кол_год']),
                                номера=a['номера'], задачи=r.get('Номера задач #', ''),
                                ш=r.get('Широта', ''), д=r.get('Долгота', ''),
                                Статус=status(a['посл']), посл=a['посл'],
                                загр=r.get('Дата загрузки', ''), новый=False))
            else:
                # несматченный адрес: жалоб 0, старая дата сохраняется (вариант А)
                out.append(dict(Улица=m['Улица'], Адрес=r['Адрес'], ОМСУ=m['ОМСУ'], кол=0, номера='',
                                задачи=r.get('Номера задач #', ''),
                                ш=r.get('Широта', ''), д=r.get('Долгота', ''),
                                Статус='Нет жалоб последние 30 дней',
                                посл=r.get('Дата последней жалобы', pd.NaT),
                                загр=r.get('Дата загрузки', ''), новый=False))
        # новые адреса на существующей улице
        if len(ag):
            for ak, a in ag.iterrows():
                if ak in used:
                    continue
                out.append(dict(Улица=m['Улица'], Адрес=a['дисп'], ОМСУ=m['ОМСУ'], кол=int(a['кол_год']),
                                номера=a['номера'], задачи=m['задачи'], ш=m['ш'], д=m['д'],
                                Статус='Новый системный', посл=a['посл'], загр=new_period, новый=True))
        cnt = len(out)
        for o in out:
            o['адресов'] = cnt
        rows += out
    adr = pd.DataFrame(rows)

    # ----- новые системные улицы -----
    wn = df[(df['дата'] >= pd.Timestamp(win_start)) & (df['дата'] < pd.Timestamp(ref_date) + timedelta(days=1))
            & df['ул_ключ'].notna() & ~df['sk'].apply(bool).astype(bool)]
    newrows = []
    for (fk, oms_key), grp in wn.groupby(['ул_ключ', 'омсу_н']):
        days = grp['дата'].dt.normalize().nunique()
        if len(grp) >= MIN_COMPLAINTS and days >= MIN_DAYS and is_real_street(grp['ул_дисп'].iloc[0]):
            disp = grp['ул_дисп'].iloc[0]; oms = grp['омсу_src'].iloc[0]
            ag = grp.groupby('а_ключ').apply(aggregate, include_groups=False)
            for ak, a in ag.iterrows():
                newrows.append(dict(Улица=disp, Адрес=a['дисп'], ОМСУ=oms, кол=int(a['кол']), кол_год=int(a['кол_год']),
                                    номера=a['номера'], задачи='', ш='', д='',
                                    Статус=status(a['посл']), посл=a['посл'],
                                    загр=new_period, новый=True, адресов=len(ag)))
    nw = pd.DataFrame(newrows)
    if len(nw):
        all_streets = pd.concat([s_str[['Улица', 'ОМСУ']], nw[['Улица', 'ОМСУ']]])
        if (all_streets.groupby('Улица')['ОМСУ'].nunique() > 1).any():
            raise ValueError('Одинаковые адреса улиц относятся к разным ОМСУ. Уточните полный адрес в источнике.')

    # ----- агрегат до улиц (Sheet1) -----
    def join_tasks(s):
        seen = []
        for v in s.dropna():
            for t in re.split(r'[,;\s]+', str(v)):
                t = t.strip()
                if t and t.lower() != 'nan' and t not in seen:
                    seen.append(t)
        return ', '.join(seen)

    g = adr.groupby('Улица', sort=False).agg(
        ОМСУ=('ОМСУ', 'first'), кол=('кол', 'sum'), адресов=('Адрес', 'count'),
        посл=('посл', 'max'), задачи=('задачи', join_tasks)).reset_index()
    g['Статус'] = g['посл'].apply(status)
    g['k'] = g['Улица'].apply(norm)

    # --- Правило возврата в статус: улица, которая на прошлой неделе была
    # "Нет жалоб последние 30 дней", возвращается в "Жалобы продолжаются"
    # только если в окне >= 2 жалоб в >= 2 разных календарных дня.
    # Уже активные улицы (были "продолжаются"/"Новый системный") под правило не попадают.
    prod_lbl = 'Жалобы продолжаются'
    net_lbl = 'Нет жалоб последние 30 дней'
    old_status = dict(zip(s_str['ул_ключ'], s_str['Статус'])) if 'Статус' in s_str.columns else {}
    # жалобы окна по каждой улице: количество и число разных дней
    win_mask = (df['дата'] >= pd.Timestamp(win_start)) & (df['дата'] < pd.Timestamp(ref_date) + timedelta(days=1))
    dwin = df[win_mask].copy()
    dwin['дом_н'] = dwin['дом'].apply(lambda h: house_key(h) if pd.notna(h) else '')
    # для каждой улицы определяем, выполнен ли порог возврата:
    #  - если у жалоб указан дом → нужен ДОМ с >= 2 жалобами в >= 2 разных дня;
    #  - жалобы без дома считаются по улице (>= 2 в >= 2 дня).
    okno_return = {}
    for skk in g['k']:
        sub = dwin[dwin['sk'].apply(lambda s: bool(s) and skk in s).astype(bool)]
        ok = False
        if len(sub):
            with_h = sub[sub['дом_н'] != '']
            for _dom, gr in with_h.groupby('дом_н'):
                if len(gr) >= 2 and gr['дата'].dt.normalize().nunique() >= 2:
                    ok = True
                    break
            if not ok:
                no_h = sub[sub['дом_н'] == '']
                if len(no_h) >= 2 and no_h['дата'].dt.normalize().nunique() >= 2:
                    ok = True
        okno_return[skk] = ok
    def adjust_status(row):
        st = row['Статус']
        if st != prod_lbl:
            return st
        prev = old_status.get(row['k'])
        prev_active = prev in (prod_lbl, 'Новый системный')
        if prev_active:
            return st  # уже была активна — достаточно любой жалобы в окне
        # была "нет жалоб" (или новая улица) → нужен порог возврата по адресу/улице
        if okno_return.get(row['k'], False):
            return prod_lbl
        return net_lbl
    g['Статус'] = g.apply(adjust_status, axis=1)
    coord = s_str.set_index('ул_ключ')
    lat_col = 'широта' if 'широта' in s_str.columns else ('Широта' if 'Широта' in s_str.columns else None)
    lon_col = 'долгота' if 'долгота' in s_str.columns else ('Долгота' if 'Долгота' in s_str.columns else None)
    g['широта'] = g['k'].map(lambda k: coord[lat_col].get(k, '') if lat_col else '')
    g['долгота'] = g['k'].map(lambda k: coord[lon_col].get(k, '') if lon_col else '')
    zg_col = 'дата загрузки' if 'дата загрузки' in s_str.columns else None
    zg = coord[zg_col].to_dict() if zg_col else {}
    g['загр'] = g['k'].map(lambda k: zg.get(k, ''))

    stats = {
        'unmatched': unmatched,
        'new_period': new_period,
        'win_start': win_start,
        'ref_date': ref_date,
        'warning': getattr(load_systemka, 'warning', ''),
    }
    return adr, nw, g, s_str, stats

# ---------------------------------------------------------------------------
# Запись Excel

def write_excel(adr, nw, g, s_str, ref_date, out_path):
    F = lambda **k: Font(name='Arial', size=10, **k)
    HDR = PatternFill('solid', start_color='1F4E79')
    NEW = PatternFill('solid', start_color='FFE699')
    REDF = PatternFill('solid', start_color='F8CBAD')
    NSTR = PatternFill('solid', start_color='FCE4D6')
    thin = Border(*[Side(style='thin', color='BFBFBF')] * 4)

    def hdr(ws, H, W):
        for i, (h, w) in enumerate(zip(H, W), 1):
            c = ws.cell(1, i, h)
            c.font = F(bold=True, color='FFFFFF'); c.fill = HDR
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = thin
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 30

    def put(ws, r, vals, fill=None, scol=None, wrap=()):
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            if isinstance(v, str):
                c.data_type = 's'
            c.font = F(); c.border = thin
            c.alignment = Alignment(vertical='top', wrap_text=(i in wrap))
            if isinstance(v, pd.Timestamp):
                c.number_format = 'DD.MM.YYYY'
            if fill:
                c.fill = fill
        if scol:
            sc = ws.cell(r, scol)
            if sc.value == 'Жалобы продолжаются':
                sc.fill = REDF; sc.font = F(bold=True)
            elif sc.value == 'Новый системный':
                sc.fill = NEW; sc.font = F(bold=True)

    date_str = ref_date.strftime('%d.%m.%Y')
    win_str = f"{(ref_date - timedelta(days=WINDOW_DAYS - 1)).strftime('%d.%m')}–{ref_date.strftime('%d.%m')}"
    new_period = f"{(ref_date - timedelta(days=WINDOW_DAYS - 1)).strftime('%d.%m')}.-{ref_date.strftime('%d.%m.%Y')}"
    has_week_flag = 'добавлено на текущей неделе' in s_str.columns

    wb = Workbook()

    # ---- Sheet1 ----
    ws = wb.active; ws.title = 'Sheet1'
    H = ['Улица', 'ОМСУ', f'Жалоб за {ref_date.year}', 'Номера задач #', 'широта', 'долгота',
         f'Статус на {date_str}', 'Категория', 'дата последней жалобы', 'дата загрузки']
    W = [58, 18, 11, 40, 10, 10, 24, 18, 16, 16]
    if has_week_flag:
        H.append('добавлено на текущей неделе'); W.append(14)
    hdr(ws, H, W)
    r = 2
    for _, x in g.sort_values(['ОМСУ', 'Улица']).iterrows():
        vals = [x['Улица'], x['ОМСУ'], int(x['кол']), x['задачи'], x['широта'], x['долгота'],
                x['Статус'], 'Системный адрес',
                x['посл'] if pd.notna(x['посл']) else '', x['загр'] if pd.notna(x['загр']) else '']
        if has_week_flag:
            vals.append(0)
        put(ws, r, vals, scol=7, wrap=(1, 4)); r += 1
    if len(nw):
        ga = nw.groupby(['ОМСУ', 'Улица']).agg(кол=('кол_год', 'sum'), посл=('посл', 'max')).reset_index()
        for _, x in ga.sort_values(['ОМСУ', 'Улица']).iterrows():
            vals = [x['Улица'], x['ОМСУ'], int(x['кол']), '', '', '', 'Жалобы продолжаются',
                    'Новая системная улица', x['посл'], new_period]
            if has_week_flag:
                vals.append(1)
            put(ws, r, vals, fill=NSTR, scol=7, wrap=(1, 4)); r += 1
    ws.auto_filter.ref = f'A1:{get_column_letter(len(H))}{r-1}'

    # ---- Новые системные улицы ----
    ws3 = wb.create_sheet('Новые системные улицы')
    H3 = ['Улица', 'Адрес', 'ОМСУ', f'Жалоб ({win_str})', 'Количество адресов',
          'Номера жалоб', 'Статус', 'Дата последней жалобы', 'Дата загрузки']
    hdr(ws3, H3, [56, 28, 16, 14, 11, 46, 22, 16, 16])
    r = 2
    if len(nw):
        for _, x in nw.sort_values(['ОМСУ', 'Улица', 'кол'], ascending=[True, True, False]).iterrows():
            put(ws3, r, [x['Улица'], x['Адрес'], x['ОМСУ'], int(x['кол']), int(x['адресов']),
                         x['номера'], x['Статус'], x['посл'], x['загр']],
                fill=NSTR, scol=7, wrap=(1, 2, 6)); r += 1
    ws3.auto_filter.ref = f'A1:I{max(r-1,1)}'

    # ---- Лист1 (адреса) ----
    ws2 = wb.create_sheet('Лист1')
    H2 = ['Улица', 'Адрес', 'ОМСУ', f'Жалоб за {ref_date.year}', 'Количество адресов', 'Номера жалоб',
          'Номера задач #', 'Широта', 'Долгота', 'Статус', 'Дата последней жалобы', 'Дата загрузки']
    hdr(ws2, H2, [52, 26, 16, 11, 11, 46, 18, 10, 10, 22, 15, 17])
    allr = pd.concat([adr.assign(_nw=False), nw.assign(_nw=True)], ignore_index=True) if len(nw) \
        else adr.assign(_nw=False)
    allr = allr.sort_values(['ОМСУ', 'Улица', 'кол'], ascending=[True, True, False])
    r = 2
    for _, x in allr.iterrows():
        fill = NSTR if x.get('_nw') else (NEW if x['новый'] else None)
        put(ws2, r, [x['Улица'], x['Адрес'], x['ОМСУ'], int(x['кол_год'] if x.get('_nw') else x['кол']), int(x['адресов']),
                     x['номера'] if pd.notna(x['номера']) else '',
                     x['задачи'] if pd.notna(x['задачи']) else '',
                     x['ш'] if x['ш'] != '' else '', x['д'] if x['д'] != '' else '',
                     x['Статус'], x['посл'] if pd.notna(x['посл']) else '', x['загр']],
            fill=fill, scol=10, wrap=(1, 2, 6, 7)); r += 1
    ws2.auto_filter.ref = f'A1:L{r-1}'

    wb.save(out_path)

# ---------------------------------------------------------------------------
# СМС-сводка

def sms_summary(g, nw, s_str):
    prod = 'Жалобы продолжаются'
    net = 'Нет жалоб последние 30 дней'
    old = dict(zip(s_str['ул_ключ'], s_str['Статус'])) if 'Статус' in s_str.columns else {}
    g = g.copy()
    g['old'] = g['k'].map(old).replace({'Новый системный': prod})
    to_net = int(((g['old'] == prod) & (g['Статус'] == net)).sum())
    to_prod = int(((g['old'] == net) & (g['Статус'] == prod)).sum())
    stay_prod = int(((g['old'] == prod) & (g['Статус'] == prod)).sum())
    stay_net = int(((g['old'] == net) & (g['Статус'] == net)).sum())
    new_streets = int(nw['Улица'].nunique()) if len(nw) else 0
    existing = len(g)
    lines = [
        f"Системные адреса: всего на контроле {existing + new_streets}, из них:",
        f"- Перешли в разряд нет жалоб: {to_net}",
        f"- Статус системный вернулся: {to_prod}",
        f"- Остались «продолжаются»: {stay_prod}",
        f"- Остались «нет жалоб»: {stay_net}",
        f"- Новые системные: {new_streets}",
        f"(проверка: {to_net}+{to_prod}+{stay_prod}+{stay_net}={existing}, +{new_streets} = {existing + new_streets})",
    ]
    return '\n'.join(lines)

# ---------------------------------------------------------------------------
# Дополнительные выгрузки

def returned_streets(g, s_str):
    """Улицы, вернувшиеся из 'нет жалоб' в 'Жалобы продолжаются'.
    Возвращает (DataFrame, строка_номеров_задач_через_запятую_без_пробелов)."""
    prod = 'Жалобы продолжаются'
    net = 'Нет жалоб последние 30 дней'
    old = dict(zip(s_str['ул_ключ'], s_str['Статус'])) if 'Статус' in s_str.columns else {}
    gg = g.copy()
    gg['old'] = gg['k'].map(old).replace({'Новый системный': prod})
    back = gg[(gg['old'] == net) & (gg['Статус'] == prod)].copy()
    if len(back):
        back = back.sort_values(['ОМСУ', 'Улица'])
    tasks = []
    for v in back['задачи']:
        if pd.isna(v):
            continue
        for t in re.split(r'[,;\s]+', str(v)):
            t = t.strip()
            if t and t.lower() != 'nan' and t not in tasks:
                tasks.append(t)
    return back, ','.join(tasks)


def write_returned_and_new(g, nw, s_str, ref_date, out_path):
    """Excel с двумя листами: вернувшиеся улицы (+задачи) и новые системные улицы (+номера жалоб)."""
    from openpyxl import Workbook
    win_start = ref_date - timedelta(days=WINDOW_DAYS - 1)
    win_str = f"{win_start.strftime('%d.%m')}–{ref_date.strftime('%d.%m')}"

    F = lambda **k: Font(name='Arial', size=10, **k)
    HDR = PatternFill('solid', start_color='1F4E79')
    NSTR = PatternFill('solid', start_color='FCE4D6')
    BACKF = PatternFill('solid', start_color='F8CBAD')
    thin = Border(*[Side(style='thin', color='BFBFBF')] * 4)

    def hdr(ws, H, W):
        for i, (h, w) in enumerate(zip(H, W), 1):
            c = ws.cell(1, i, h)
            c.font = F(bold=True, color='FFFFFF'); c.fill = HDR
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = thin
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 30

    def put(ws, r, vals, fill, wrap=()):
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            if isinstance(v, str):
                c.data_type = 's'
            c.font = F(); c.border = thin; c.fill = fill
            c.alignment = Alignment(vertical='top', wrap_text=(i in wrap))
            if isinstance(v, pd.Timestamp):
                c.number_format = 'DD.MM.YYYY'

    wb = Workbook()

    # --- лист 1: вернувшиеся улицы ---
    back, tasks_line = returned_streets(g, s_str)
    ws = wb.active; ws.title = 'Вернувшиеся улицы'
    hdr(ws, ['ОМСУ', 'Улица', f'Жалоб за {ref_date.year}', 'Дата последней жалобы', 'Номера задач #'],
        [18, 66, 12, 16, 34])
    r = 2
    for _, x in back.iterrows():
        tks = ', '.join(dict.fromkeys(
            [t for t in re.split(r'[,;\s]+', str(x['задачи'])) if t and t.lower() != 'nan']))
        put(ws, r, [x['ОМСУ'], x['Улица'], int(x['кол']),
                    x['посл'] if pd.notna(x['посл']) else '', tks], BACKF, wrap=(2, 5))
        r += 1
    if r > 2:
        ws.auto_filter.ref = f'A1:E{r-1}'
    # сводная строка со всеми задачами
    r += 1
    c = ws.cell(r, 1, 'Все номера задач:'); c.font = F(bold=True)
    c2 = ws.cell(r, 2, tasks_line); c2.data_type = 's'; c2.font = F(); c2.alignment = Alignment(wrap_text=True, vertical='top')

    # --- лист 2: новые системные улицы ---
    ws2 = wb.create_sheet('Новые системные улицы')
    hdr(ws2, ['Улица', 'ОМСУ', f'Жалоб ({win_str})', 'Номера жалоб'], [72, 18, 14, 58])
    r = 2
    if len(nw):
        ga = nw.groupby(['Улица', 'ОМСУ'], sort=False).agg(
            кол=('кол', 'sum'),
            номера=('номера', lambda s: ', '.join(dict.fromkeys(
                [t.strip() for v in s.dropna() for t in str(v).split(',') if t.strip()])))
        ).reset_index().sort_values(['ОМСУ', 'Улица'])
        for _, x in ga.iterrows():
            put(ws2, r, [x['Улица'], x['ОМСУ'], int(x['кол']), x['номера']], NSTR, wrap=(1, 4))
            r += 1
    if r > 2:
        ws2.auto_filter.ref = f'A1:D{r-1}'

    wb.save(out_path)
    return tasks_line
