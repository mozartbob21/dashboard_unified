# -*- coding: utf-8 -*-
"""Системные адреса: reviewed local engine and protected persisted registries.

The caller enforces account rights and serializes jobs/imports. Portal data is
fetched only through the existing МинЖКХ credential client; no env-file access,
runtime exec, uploaded scripts or statically served workbooks are used here.
"""
import base64
import concurrent.futures
import datetime
import hashlib
import io
import json
import os
from pathlib import Path
import re
import tempfile
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook

from .context import DATA, today
from . import dataset, system_engine
from .client import Pentaho

DIRECTION_WATER = "6"
OUT = Path(DATA) / "sys_out"
REG_DIR = Path(DATA) / "sys_registry"
MAX_REGISTRY_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_EXPORT_BYTES = 20 * 1024 * 1024
MAX_REGISTRY_ROWS = 100_000
MAX_REGISTRY_COLUMNS = 60
PROD = "Жалобы продолжаются"
NET = "Нет жалоб последние 30 дней"


def _report_date(value):
    try:
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', str(value)):
            raise ValueError
        day = datetime.date.fromisoformat(str(value))
    except (TypeError, ValueError):
        raise ValueError('Укажите дату отчёта в формате ГГГГ-ММ-ДД.') from None
    if day.year < 2000 or day > today():
        raise ValueError('Дата отчёта должна быть между 01.01.2000 и сегодняшним днём.')
    return day


def _atomic_write(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix='.' + path.name, dir=path.parent)
    try:
        with os.fdopen(fd, 'wb') as target:
            target.write(data)
            target.flush()
            os.fsync(target.fileno())
        os.replace(tmp, path)
    finally:
        Path(tmp).unlink(missing_ok=True)


def registries():
    """Valid local registry dates, oldest first; ignored files are never loaded."""
    if not REG_DIR.is_dir():
        return []
    dates = []
    for path in REG_DIR.iterdir():
        if not path.is_file() or path.is_symlink() or not re.fullmatch(r'\d{4}-\d{2}-\d{2}\.xlsx', path.name):
            continue
        try:
            dates.append(datetime.date.fromisoformat(path.stem))
        except ValueError:
            continue
    return sorted(dates)


def previous(ref_d):
    ref_d = _report_date(ref_d)
    saved = registries()
    earlier = [day for day in saved if day < ref_d]
    if not earlier:
        first = saved[0].strftime('%d.%m.%Y') if saved else '—'
        raise ValueError('Нет сохранённого реестра раньше %s (самый ранний — %s). '
                         'Загрузите исходный реестр в этом разделе.'
                         % (ref_d.strftime('%d.%m.%Y'), first))
    return earlier[-1]


def _xlsx_rows(payload):
    if not isinstance(payload, bytes) or not payload or len(payload) > MAX_REGISTRY_BYTES:
        raise ValueError('Загрузите непустой файл XLSX размером не более 20 МБ.')
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if len(entries) > 2000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise ValueError('Файл XLSX содержит слишком большой объём данных.')
            if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
                raise ValueError('Нужен файл Excel в формате XLSX.')
            if any('vbaproject' in name.lower() or '/externallinks/' in name.lower() for name in names):
                raise ValueError('Реестр не должен содержать макросы или внешние связи.')
            if any(entry.flag_bits & 1 for entry in entries):
                raise ValueError('Защищённые паролем XLSX не поддерживаются.')
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        try:
            selected = ['Sheet1' if 'Sheet1' in workbook.sheetnames else workbook.sheetnames[0]]
            if 'Лист1' in workbook.sheetnames and 'Лист1' not in selected:
                selected.append('Лист1')
            result = {}
            for name in selected:
                sheet = workbook[name]
                if sheet.max_row and sheet.max_row > MAX_REGISTRY_ROWS + 1 or sheet.max_column and sheet.max_column > MAX_REGISTRY_COLUMNS:
                    raise ValueError('Реестр содержит слишком много строк или столбцов.')
                rows = []
                for number, row in enumerate(sheet.iter_rows(), 1):
                    if number > MAX_REGISTRY_ROWS + 1 or len(row) > MAX_REGISTRY_COLUMNS:
                        raise ValueError('Реестр содержит слишком много строк или столбцов.')
                    values = []
                    for cell in row:
                        if cell.data_type == 'f':
                            raise ValueError('Реестр содержит формулы. Сохраните его как значения и загрузите снова.')
                        value = cell.value
                        if isinstance(value, str) and len(value) > 32767:
                            raise ValueError('В реестре слишком длинное значение ячейки.')
                        values.append(value)
                    if any(value is not None for value in values):
                        rows.append(values)
                result['Sheet1' if name == selected[0] else name] = rows
            return result
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception:
        raise ValueError('Не удалось прочитать XLSX. Проверьте формат и целостность реестра.') from None


def prepare_registry(payload):
    """Return a validated values-only XLSX without writing to runtime storage."""
    sheets = _xlsx_rows(payload)
    street_keys = set()
    total = 0
    for name, rows in sheets.items():
        if len(rows) < 2:
            raise ValueError('Реестр должен содержать заголовки и хотя бы одну запись.')
        headers = [str(value).strip() if value is not None else '' for value in rows[0]]
        required = {'Улица', 'ОМСУ'} | ({'Адрес'} if name == 'Лист1' else set())
        if not required <= set(headers):
            raise ValueError('В листе %s нужны колонки: %s.' % (name, ', '.join(sorted(required))))
        statuses = [i for i, value in enumerate(headers) if value.startswith('Статус')]
        if name == 'Sheet1' and len(statuses) != 1:
            raise ValueError('В листе улиц нужна одна колонка «Статус» или «Статус на дату».')
        if len(set(header for header in headers if header)) != len([header for header in headers if header]):
            raise ValueError('В реестре повторяются названия колонок.')
        rows[0] = headers
        for row in rows[1:]:
            street = row[headers.index('Улица')]
            omsu = row[headers.index('ОМСУ')]
            if not isinstance(street, str) or not street.strip() or not isinstance(omsu, str) or not omsu.strip():
                raise ValueError('У каждой записи реестра должны быть заполнены улица и ОМСУ.')
            if name == 'Sheet1':
                key = system_engine.norm(street)
                if key in street_keys:
                    raise ValueError('Улицы в реестре повторяются. Укажите полный уникальный адрес улицы, включая населённый пункт.')
                street_keys.add(key)
                status = row[statuses[0]]
                if status not in {PROD, NET, 'Новый системный'}:
                    raise ValueError('В реестре неизвестный статус улицы. Используйте результат предыдущего расчёта.')
                total += 1
            elif system_engine.norm(street) not in street_keys:
                raise ValueError('Лист адресов содержит улицу, которой нет в листе Sheet1.')
    output = Workbook()
    output.remove(output.active)
    for name, rows in sheets.items():
        sheet = output.create_sheet(name)
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    buf = io.BytesIO()
    output.save(buf)
    data = buf.getvalue()
    if len(data) > MAX_REGISTRY_BYTES:
        raise ValueError('Результирующий реестр превышает 20 МБ.')
    return data


def registry_fingerprint(payload):
    """Stable content fingerprint (not ZIP timestamps or workbook metadata)."""
    value = json.dumps(_xlsx_rows(payload), ensure_ascii=False, sort_keys=True,
                       default=lambda value: value.isoformat(), allow_nan=False)
    return hashlib.sha256(value.encode('utf-8')).hexdigest()


def import_registry(payload, report_date):
    """Publish a validated initial registry; identical retries are idempotent.

    Admin/session checks and the shared job lock belong to the API caller.
    An existing different registry is never silently replaced by an upload.
    """
    day = _report_date(report_date)
    data = prepare_registry(payload)
    target = REG_DIR / (day.isoformat() + '.xlsx')
    if target.exists():
        if target.is_symlink() or target.stat().st_size > MAX_REGISTRY_BYTES:
            raise ValueError('Реестр на эту дату уже существует и не может быть заменён.')
        if registry_fingerprint(target.read_bytes()) != registry_fingerprint(data):
            raise ValueError('На эту дату уже сохранён другой реестр. Выберите правильную дату исходного реестра.')
        created = False
    else:
        _atomic_write(target, data)
        created = True
    total = len(_xlsx_rows(data)['Sheet1']) - 1
    return {'date': day.isoformat(), 'streets': total, 'name': day.isoformat() + '.xlsx', 'created': created}


def complaints(date_to, path, log=None):
    """Fetch a full rolling window; return only the report-year complaint count."""
    log = log or (lambda _s: None)
    current_day = today()
    start = min(datetime.date(date_to.year, 1, 1), date_to - datetime.timedelta(days=system_engine.WINDOW_DAYS - 1))
    portal = Pentaho()
    base = portal.periods(10, 0, 20, 11, direction=DIRECTION_WATER)
    def one(bounds):
        params = dict(base, curr_period_start='%d day' % bounds[0], curr_period_end='%d day' % bounds[1])
        cols, rows = portal.query('q_download_detalization_address', params, timeout=300)
        return pd.DataFrame(rows, columns=cols)
    with concurrent.futures.ThreadPoolExecutor(max_workers=dataset.WORKERS) as executor:
        parts = list(executor.map(one, dataset.chunks((current_day - start).days, (current_day - date_to).days)))
    if not parts:
        raise ValueError('Портал не вернул выгрузку жалоб за выбранный период.')
    frame = pd.concat(parts, ignore_index=True)
    if 'Внутренний Id' not in frame.columns:
        raise ValueError('Портал вернул выгрузку без идентификаторов обращений.')
    frame = frame.drop_duplicates('Внутренний Id')
    frame = frame.rename(columns={'номер ЕЦУР': 'Номер ЕЦУР', 'description': 'Описание',
                                   'petitioner_email': 'Почта заявителя'})
    # Engine needs these fields only. Do not persist personal contact data.
    needed = [column for column in frame.columns if any(token in str(column).lower()
              for token in ('номер ецур', 'номер в источнике', 'дата', 'тип сообщения', 'адрес', 'дом', 'омсу'))]
    if not any('дата' in str(column).lower() for column in needed):
        raise ValueError('В выгрузке отсутствует дата обращения.')
    safe = frame[needed].copy()
    safe.to_excel(path, sheet_name='Sheet1', index=False, engine='openpyxl')
    # Keep untrusted strings as literal Excel cells, never spreadsheet formulas.
    workbook = load_workbook(path, data_only=False, keep_links=False)
    try:
        for row in workbook.active:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        workbook.save(path)
    finally:
        workbook.close()
    # The January workbook also contains December for the rolling criterion.
    # UI says «за год», so reuse the engine's validated dates/type/end-day filter
    # and keep the returned counter separate from the downloaded window size.
    records = system_engine.load_complaints(path, pd.Timestamp(date_to))
    year_count = int(records['отчётный_год'].sum())
    log('Жалоб ВС/ВО за %d год по дату отчёта: %d' % (date_to.year, year_count))
    return year_count


# описание задачи RM для новой системной улицы — как в «загрузка в RM\системные.csv»
RM_DESC = (
    '<p>Необходимо провести анализ адреса и поступивших по нему жалоб, предоставить:<br><br></p>'
    '<p>1. Подтверждение, есть ли проблематика по жалобам или нет (если нет, расписать, почему '
    'данный адрес не является проблемным).<br>2. Определить быстрое решение для исключения '
    'поступления новых жалоб.<br>2.1. Обеспечить информирование жителей данного дома/улицы по '
    'принятым решениям и срокам и предоставить подтверждающие материалы.<br>3. Определить '
    'системное решение для исключения поступления новых жалоб в принципе.<br>3.1. Обеспечить '
    'информирование жителей данного дома/улицы по принятым решениям и срокам и предоставить '
    'подтверждающие материалы.<br>4. По адресу, где фиксируется наличие резонансных обращений, '
    'обязательная встреча с жителями для очного разъяснения проблематики и путей их решения.'
    '<br><br></p><p>Скачать информационный ролик по промывкам можно <a href="https://cloud.mosreg.ru/'
    's/AgI1k7IvhaMaHY7/download?path=%2F&files=WATER%20SUPPLY.mp4">по ссылке</a></p>\n'
)

def _counts(g, nw, s_str):
    """Те же счётчики, что sms_summary движка."""
    old = dict(zip(s_str["ул_ключ"], s_str["Статус"])) if "Статус" in s_str.columns else {}
    g = g.copy()
    g["old"] = g["k"].map(old).replace({"Новый системный": PROD})
    c = {
        "to_net": int(((g["old"] == PROD) & (g["Статус"] == NET)).sum()),
        "to_prod": int(((g["old"] == NET) & (g["Статус"] == PROD)).sum()),
        "stay_prod": int(((g["old"] == PROD) & (g["Статус"] == PROD)).sum()),
        "stay_net": int(((g["old"] == NET) & (g["Статус"] == NET)).sum()),
        "new": int(nw["Улица"].nunique()) if len(nw) else 0,
        "existing": len(g),
    }
    c["total"] = c["existing"] + c["new"]
    return c, g


def sms(c):
    return "\n".join([
        "Системные адреса: всего на контроле %d (+%d новых), из них:" % (c["total"], c["new"]),
        "Перешли в категорию «нет жалоб за 30 дней»: %d✅" % c["to_net"],
        "Вернули на доработку «системный»: %d" % c["to_prod"],
        "Остаются в категории «продолжаются»: %d" % c["stay_prod"],
        "Остаются в категории «нет жалоб за 30 дней»: %d" % c["stay_net"],
        "Новые системные: %d" % c["new"],
    ])


def _b64(path):
    path = Path(path)
    if path.stat().st_size > MAX_EXPORT_BYTES:
        raise ValueError('Выгрузка слишком велика. Сократите объём реестра.')
    return base64.b64encode(path.read_bytes()).decode('ascii')


def _date(v):
    return v.strftime("%d.%m.%Y") if isinstance(v, (pd.Timestamp, datetime.date)) and not pd.isna(v) else ""


def run(req, log=None):
    """req: date (ISO) — дата отсчёта. Реестр прошлой недели — последний
    сохранённый раньше этой даты; результат сохраняется как реестр на дату."""
    log = log or (lambda _s: None)
    ref_d = _report_date(req.get('date'))
    stamp = ref_d.strftime("%d.%m.%Y")
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.system-run-', dir=OUT) as temp:
        return _run(ref_d, stamp, Path(temp), log)


def _run(ref_d, stamp, work, log):

    prev_d = previous(ref_d)
    src = os.path.join(REG_DIR, prev_d.isoformat() + ".xlsx")
    src_info = {"name": "реестр на " + prev_d.strftime("%d.%m.%Y"), "date": prev_d.isoformat()}
    compl = os.path.join(work, "жалобы_ВС_ВО.xlsx")
    n_compl = complaints(ref_d, compl, log)

    ref = pd.Timestamp(ref_d)
    adr, nw, g, s_str, stats = system_engine.process(src, compl, ref)
    reg = os.path.join(work, "системные_обновлено_%s.xlsx" % stamp)
    system_engine.write_excel(adr, nw, g, s_str, ref, reg)
    extra = os.path.join(work, "вернувшиеся_и_новые_%s.xlsx" % stamp)
    tasks_line = system_engine.write_returned_and_new(g, nw, s_str, ref, extra)
    c, gg = _counts(g, nw, s_str)
    assert c["to_net"] + c["to_prod"] + c["stay_prod"] + c["stay_net"] == c["existing"], c
    # новые системные улицы: улица, ОМСУ, жалоб в окне, номера жалоб
    new_streets = []
    if len(nw):
        ga = nw.groupby(["Улица", "ОМСУ"], sort=False).agg(
            кол=("кол", "sum"), посл=("посл", "max"),
            номера=("номера", lambda s: ", ".join(dict.fromkeys(
                t.strip() for v in s.dropna() for t in str(v).split(",") if t.strip())))
        ).reset_index().sort_values(["ОМСУ", "Улица"])
        new_streets = [{"street": r["Улица"], "omsu": r["ОМСУ"], "n": int(r["кол"]),
                        "last": _date(r["посл"]), "nums": r["номера"]} for _, r in ga.iterrows()]

    def streets(mask):
        x = gg[mask].sort_values(["ОМСУ", "Улица"])
        return [{"street": r["Улица"], "omsu": r["ОМСУ"], "n": int(r["кол"]), "last": _date(r["посл"]),
                 "tasks": [t for t in re.split(r"[,;\s]+", str(r["задачи"] or "")) if t and t.lower() != "nan"]}
                for _, r in x.iterrows()]

    result = {
        "sms": sms(c), "counts": c, "source": src_info, "complaints": n_compl,
        "unmatched": [list(u) for u in stats["unmatched"]], "warning": stats.get("warning", ""),
        "new_streets": new_streets,
        "returned": streets((gg["old"] == NET) & (gg["Статус"] == PROD)),
        "to_net": streets((gg["old"] == PROD) & (gg["Статус"] == NET)),
        "tasks_line": tasks_line,
        "files": {
            "registry": {"name": os.path.basename(reg), "b64": _b64(reg)},
            "extra": {"name": os.path.basename(extra), "b64": _b64(extra)},
        },
        # CSV новых системных для RM собирает страница: контролёр и даты вводятся
        # после пересчёта; формат — как «загрузка в RM\системные.csv»
        "rm": "https://mgkh.rm.mosreg.ru", "rm_desc": RM_DESC, "stamp": stamp,
    }

    # Validate and serialize the entire response before publishing a new registry.
    # A failed query, export or size limit therefore leaves previous results intact.
    json.dumps(result, ensure_ascii=False, allow_nan=False)
    REG_DIR.mkdir(parents=True, exist_ok=True)
    destination = REG_DIR / (ref_d.isoformat() + '.xlsx')
    _atomic_write(destination, Path(reg).read_bytes())
    return result
