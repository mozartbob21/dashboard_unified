"""Merge the first sheets, matching headers, as in the user's merge_excel_cli.py."""
import io
from pathlib import Path
from .workspace import ToolError
from .documents import validate_office

OUTPUT_NAME = 'ОБЪЕДИНЕННЫЙ_РЕЗУЛЬТАТ.xlsx'


def merge(files, job):
    import pandas as pd
    from openpyxl import Workbook
    frames, sources = [], []
    for name, data in sorted(files, key=lambda pair: pair[0].casefold()):
        if name.startswith('~$') or name.casefold() == OUTPUT_NAME.casefold():
            continue
        suffix = Path(name).suffix.lower()
        if suffix == '.xlsx':
            validate_office(data)
        try:
            frame = pd.read_excel(io.BytesIO(data), sheet_name=0,
                                  engine={'.xlsx':'openpyxl','.xls':'xlrd','.xlsb':'pyxlsb'}[suffix])
        except ImportError:
            raise ToolError(f'Для {suffix} нужна библиотека ' + {'.xls':'xlrd','.xlsb':'pyxlsb'}.get(suffix, 'openpyxl') + '. Обновите зависимости.')
        except Exception:
            raise ToolError(f'Не удалось прочитать «{name}». Результат не создан, чтобы не потерять строки.')
        frame.columns = [str(c).strip() for c in frame.columns]
        if len(set(frame.columns)) != len(frame.columns):
            raise ToolError(f'В «{name}» повторяются заголовки столбцов.')
        frames.append(frame)
        sources.append((name, len(frame)))
        if sum(f.size for f in frames) > 500000:
            raise ToolError('Слишком много данных: ограничение — 500 000 ячеек.')
    if not frames:
        raise ToolError('Нет подходящих Excel-файлов. Временные файлы и предыдущий результат пропускаются.')
    combined = pd.concat(frames, ignore_index=True, sort=False).astype(object)
    combined = combined.where(pd.notna(combined), None)
    if combined.size > 500000:
        raise ToolError('После объединения получилось больше 500 000 ячеек.')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Данные'
    for row in [list(combined.columns), *combined.itertuples(index=False, name=None)]:
        sheet.append(list(row))
        for cell in sheet[sheet.max_row]:
            # A text value beginning with '=' must never become a new Excel formula.
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    manifest = workbook.create_sheet('Источники')
    manifest.append(['Файл', 'Строк'])
    for name, count in sources:
        manifest.append([name, count])
        manifest.cell(manifest.max_row, 1).data_type = 's'
    workbook.save(job / OUTPUT_NAME)
    return {'file': OUTPUT_NAME, 'message': f'Объединено файлов: {len(frames)}. Строк: {len(combined)}.', 'rows':len(combined), 'sources':len(frames)}
