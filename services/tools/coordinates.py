"""Coordinate grouping adapted from the owner's «кодер объектов.py».

Complete linkage for new points, stable medoid-based codes, and a private,
atomic registry. Processing is local; no geocoding or map API calls.
"""
import html
import io
import json
import math
import os
import re
import threading
import time
from datetime import datetime
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from .documents import validate_office
from .workspace import ToolError

LAT_MIN, LAT_MAX = 54.0, 57.3
LON_MIN, LON_MAX = 34.8, 40.6
CODE_DIGITS = 4
USE_REGISTRY = True
PLACEHOLDER_MIN_ROWS = 5
PLACEHOLDER_MIN_SHARE = 0.15
PLACEHOLDER_ROWS_NO_OMSU = 20
REG_COLS = ['Код объекта', 'Префикс', 'Номер', 'Широта', 'Долгота', 'Тип точки', 'Файл', 'Добавлен']
_REGISTRY_LOCK = threading.Lock()


def _mode(series):
    values = series.dropna().mode()
    return values.iloc[0] if len(values) else ''


def _clean(s):
    s = html.unescape(str(s))
    s = re.sub(r'\(.*?\)', ' ', s)
    return s.strip()


def _variants(s):
    """Прочтения одной координаты: [(градусы, 'как записано'|'десятичные пары')]."""
    groups = re.findall(r'\d+', s)
    seps = [x.strip() for x in re.findall(r'\d+(\D*)', s)[:-1]]
    n, out = len(groups), []
    dec = lambda sep: sep in ('.', ',')  # noqa: E731
    if n == 1 and len(groups[0]) <= 3:
        out.append((float(groups[0]), 'native'))
    elif n == 2:
        if dec(seps[0]):
            out.append((float(f'{groups[0]}.{groups[1]}'), 'native'))
        elif int(groups[1]) < 60:
            out.append((int(groups[0]) + int(groups[1]) / 60, 'native'))
    elif n == 3:
        d, m, sec = groups
        if not dec(seps[0]) and dec(seps[1]):                       # 55°30.931'
            mm = float(f'{m}.{sec}')
            if mm < 60:
                out.append((int(d) + mm / 60, 'native'))
        else:
            if int(m) <= 60 and int(sec) <= 60:
                out.append((int(d) + int(m) / 60 + int(sec) / 3600, 'native'))
            if len(m) == 2 and len(sec) >= 2:                       # 55-93-28 = 55.9328
                out.append((float(f'{d}.{m}{sec}'), 'pairs'))
    elif n == 4:
        d, m = int(groups[0]), int(groups[1])
        sec = float(f'{groups[2]}.{groups[3]}')
        if m <= 60 and sec <= 60:
            out.append((d + m / 60 + sec / 3600, 'native'))
    elif n == 1 and len(groups[0]) >= 6:                            # 55545503 -> 55.545503
        out.append((float(groups[0][:2] + '.' + groups[0][2:]), 'pairs'))
    return out


def _in_box(la, lo):
    return LAT_MIN <= la <= LAT_MAX and LON_MIN <= lo <= LON_MAX


def _split(raw):
    s = _clean(raw)
    pair = re.fullmatch(r'([+-]?\d{1,3}(?:[.,]\d+)?)\s*[;,\s]\s*([+-]?\d{1,3}(?:[.,]\d+)?)', s)
    if pair:
        return [pair[1], pair[2]]
    parts = [p.strip() for p in re.split(r';|,\s+(?=\d)|(?<=[NnEeСсВв"”])\s+(?=\d)', s) if re.search(r'\d', p)]
    multi = len(parts) > 1 and (len(re.findall(r'\d+', s)) > 4 or re.fullmatch(r'[\d.,\s]+', s))
    return parts if multi else ([s] if re.search(r'\d', s) else [])


def parse_coords(lat_raw, lon_raw=None):
    """-> (широта, долгота, примечание) или (None, None, причина)."""
    def num(x):
        return isinstance(x, (int, float)) and not (isinstance(x, float) and math.isnan(x))

    if any(re.search(r'^\s*-|[SWЮЗ]\s*$', str(v), re.I) for v in (lat_raw, lon_raw) if v is not None):
        return None, None, 'вне допустимой рамки'
    if num(lat_raw) and num(lon_raw):
        la, lo = float(lat_raw), float(lon_raw)
        if _in_box(la, lo):
            return la, lo, ''
        if _in_box(lo, la):
            return lo, la, 'широта и долгота были перепутаны'
        return None, None, f'вне допустимой рамки ({la}, {lo})'
    lat_s = '' if lat_raw is None or (isinstance(lat_raw, float) and math.isnan(lat_raw)) else str(lat_raw)
    lon_s = '' if lon_raw is None or (isinstance(lon_raw, float) and math.isnan(lon_raw)) else str(lon_raw)
    lp, op = ([_clean(lat_s)], [_clean(lon_s)]) if lon_s and _clean(lat_s) != _clean(lon_s) else (_split(lat_s), _split(lon_s))
    if not lp:
        return None, None, 'нет координат'
    if not op or _clean(lat_s) == _clean(lon_s):
        if len(lp) < 2:
            return None, None, 'нет второй координаты'
        a_s, b_s = lp[0], lp[1]
    else:
        a_s, b_s = lp[0], op[0]
    if any(re.search(r'^\s*-|[SWЮЗ]\s*$', v, re.I) for v in (a_s, b_s)):
        return None, None, 'вне допустимой рамки'
    best = None
    for a, ha in _variants(a_s):
        for b, hb in _variants(b_s):
            for la, lo, swapped in ((a, b, False), (b, a, True)):
                if _in_box(la, lo):
                    rank = (ha != 'native') + (hb != 'native') + (0.5 if swapped else 0)
                    if best is None or rank < best[0]:
                        best = (rank, la, lo, swapped, 'pairs' in (ha, hb))
    if best is None:
        return None, None, 'не распознаны или вне рамки'
    _, la, lo, swapped, pairs = best
    note = []
    if swapped:
        note.append('широта и долгота были перепутаны')
    if pairs:
        note.append('прочитано как десятичные градусы, разбитые на пары цифр — проверить')
    return la, lo, '; '.join(note)


def haversine(lat1, lon1, lat2, lon2):
    """Расстояние в метрах; принимает числа или массивы numpy."""
    p1, p2 = np.radians(lat1), np.radians(lat2)
    dp, dl = p2 - p1, np.radians(lon2 - lon1)
    h = np.sin(dp / 2) ** 2 + np.cos(p1) * np.cos(p2) * np.sin(dl / 2) ** 2
    return 2 * 6371000.0 * np.arcsin(np.sqrt(np.clip(h, 0, 1)))


class Grid:
    """Сетка для быстрого поиска соседей в радиусе R."""

    def __init__(self, lats, lons, r, lat0):
        self.cell = r * 1.2
        self.kx = 111320.0 * math.cos(math.radians(lat0))
        self.cells = {}
        for i, (la, lo) in enumerate(zip(lats, lons)):
            self.cells.setdefault(self.key(la, lo), []).append(i)

    def key(self, la, lo):
        return int(math.floor(lo * self.kx / self.cell)), int(math.floor(la * 110574.0 / self.cell))

    def near(self, la, lo):
        cx, cy = self.key(la, lo)
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                yield from self.cells.get((cx + dx, cy + dy), ())


def complete_linkage(D, r):
    """Иерархическая кластеризация с полной связью: в группе все пары не дальше r."""
    k = len(D)
    if k == 1:
        return [[0]]
    M = D.astype(float).copy()
    np.fill_diagonal(M, np.inf)
    members = [[i] for i in range(k)]
    alive = np.ones(k, bool)
    while True:
        idx = int(np.argmin(M))
        a, b = divmod(idx, k)
        if not np.isfinite(M[a, b]) or M[a, b] > r:
            break
        a, b = min(a, b), max(a, b)
        members[a] += members[b]
        members[b] = None
        alive[b] = False
        row = np.maximum(M[a], M[b])
        row[~alive] = np.inf
        row[a] = np.inf
        M[a, :] = row
        M[:, a] = row
        M[b, :] = np.inf
        M[:, b] = np.inf
    return [m for m in members if m is not None]


def find_col(columns, keys, exclude=()):
    for c in columns:
        n = str(c).strip().lower()
        if any(k in n for k in keys) and not any(e in n for e in exclude):
            return c
    return None


def prefix_for(path):
    m = re.search(r'тип\s+([A-Za-zА-Яа-яЁё0-9]+)', os.path.basename(path), re.I)
    return m.group(1).upper() if m else 'OBJ'


def style(path):
    wb = load_workbook(path)
    warn = PatternFill('solid', fgColor='FFF2CC')
    hub = PatternFill('solid', fgColor='F8CBAD')
    for ws in wb.worksheets:
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True)
        for col in ws.columns:
            wd = max(len(str(c.value)) if c.value is not None else 0 for c in col[:500])
            ws.column_dimensions[col[0].column_letter].width = min(max(wd + 2, 10), 60)
        hdr = [c.value for c in ws[1]]
        kt = hdr.index('Тип точки') if 'Тип точки' in hdr else None
        kn = next((hdr.index(h) for h in ('Примечание', 'Примечание по координатам') if h in hdr), None)
        if kt is None and kn is None:
            continue
        for row in ws.iter_rows(min_row=2):
            fill = None
            if kt is not None and row[kt].value == 'условная точка округа':
                fill = hub
            elif kn is not None and row[kn].value:
                fill = warn
            if fill:
                for c in row:
                    c.fill = fill
    for ws in wb.worksheets:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
    wb.save(path)


def _process(src, registry, name, radius, prefix, job):
    deadline = time.monotonic() + 45
    src = src.dropna(how='all')
    cols = list(src.columns)
    lat_c = find_col(cols, ('широт', 'lat'))
    lon_c = find_col(cols, ('долгот', 'lon', 'lng'))
    both_c = None if (lat_c and lon_c) else find_col(cols, ('координат', 'coord'))
    id_c = find_col(cols, ('id', 'пин', 'код', 'номер', '№'))
    omsu_c = find_col(cols, ('омсу', 'округ', 'муницип'))
    rso_c = find_col(cols, ('рсо', 'исполнит', 'организац'))
    if not ((lat_c and lon_c) or both_c):
        raise ToolError('Нужны столбцы «Широта» и «Долгота» либо один столбец «Координаты» в первой строке.')

    prefix = prefix or prefix_for(name)

    # 1. координаты
    parsed = [parse_coords(r[lat_c], r[lon_c]) if lat_c and lon_c else parse_coords(r[both_c])
              for _, r in src.iterrows()]
    src = src.reset_index(drop=True)
    src['_lat'] = [p[0] for p in parsed]
    src['_lon'] = [p[1] for p in parsed]
    src['_note'] = [p[2] for p in parsed]
    ok = src['_lat'].notna()
    bad = src[~ok]
    d = src[ok].copy()
    if d.empty:
        raise ToolError('Нет распознанных координат в Московской области. Проверьте формат и выбранный файл.')

    # 2. уникальные точки
    d['_lat'] = d['_lat'].astype(float).round(7)
    d['_lon'] = d['_lon'].astype(float).round(7)
    pts = d.groupby(['_lat', '_lon']).size().reset_index(name='w')          # отсортированы по широте
    P_lat, P_lon, W = pts['_lat'].to_numpy(), pts['_lon'].to_numpy(), pts['w'].to_numpy()
    n = len(pts)
    if n > 8000:
        raise ToolError('Больше 8 000 разных точек. Разделите файл по округам.')
    pidx = {(a, b): i for i, (a, b) in enumerate(zip(P_lat, P_lon))}
    d['_p'] = [pidx[(a, b)] for a, b in zip(d['_lat'], d['_lon'])]

    # 3. привязка к реестру
    reg = registry.copy()
    regp = reg[reg['Префикс'].astype(str) == prefix].reset_index(drop=True)
    lat0 = float(np.mean(P_lat))
    assigned = {}                                                            # точка -> индекс в regp
    if len(regp):
        rg = Grid(regp['Широта'].astype(float), regp['Долгота'].astype(float), radius, lat0)
        R_lat, R_lon = regp['Широта'].astype(float).to_numpy(), regp['Долгота'].astype(float).to_numpy()
        for i in range(n):
            cand = list(rg.near(P_lat[i], P_lon[i]))
            if cand:
                dist = haversine(P_lat[i], P_lon[i], R_lat[cand], R_lon[cand])
                j = int(np.argmin(dist))
                if dist[j] <= radius:
                    assigned[i] = cand[j]

    # 4. новые точки: связные компоненты -> полная связь
    free = [i for i in range(n) if i not in assigned]
    parent = {i: i for i in free}

    def root(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    if free:
        g = Grid(P_lat[free], P_lon[free], radius, lat0)
        for a_pos, i in enumerate(free):
            if time.monotonic() > deadline:
                raise ToolError('Слишком сложная выборка. Разделите файл по округам.')
            cand = [free[c] for c in g.near(P_lat[i], P_lon[i]) if free[c] > i]
            if cand:
                dist = haversine(P_lat[i], P_lon[i], P_lat[cand], P_lon[cand])
                for j, dd in zip(cand, dist):
                    if dd <= radius:
                        parent[root(i)] = root(j)
    comps = {}
    for i in free:
        comps.setdefault(root(i), []).append(i)
    new_groups = []
    for members in comps.values():
        members.sort()
        if len(members) == 1:
            new_groups.append(members)
            continue
        if len(members) > 1000:
            raise ToolError('Больше 1 000 связанных точек. Уменьшите радиус или разделите файл.')
        la, lo = P_lat[members], P_lon[members]
        D = haversine(la[:, None], lo[:, None], la[None, :], lo[None, :])
        for cl in complete_linkage(D, radius):
            new_groups.append(sorted(members[c] for c in cl))

    # 5. группы -> объекты
    objects = []            # dict: members, lat, lon, reg_row (или None)
    by_reg = {}
    for i, j in assigned.items():
        by_reg.setdefault(j, []).append(i)
    for j, members in by_reg.items():
        objects.append({'members': sorted(members), 'lat': float(regp.at[j, 'Широта']),
                        'lon': float(regp.at[j, 'Долгота']), 'reg': j})
    for members in new_groups:
        la, lo = P_lat[members], P_lon[members]
        D = haversine(la[:, None], lo[:, None], la[None, :], lo[None, :])
        med = members[int(np.argmin((D * W[members][None, :]).sum(1)))]         # медоид с весом записей
        objects.append({'members': members, 'lat': float(P_lat[med]), 'lon': float(P_lon[med]), 'reg': None})

    point_obj = {}
    for k, o in enumerate(objects):
        for i in o['members']:
            point_obj[i] = k
    d['_o'] = d['_p'].map(point_obj)

    # условные точки
    hub_pts = set()
    if omsu_c:
        st = d.groupby('_p').agg(nn=(omsu_c, 'size'), om=(omsu_c, lambda s: s.pipe(_mode))).reset_index()
        st['share'] = st['nn'] / st['om'].map(d[omsu_c].value_counts())
        st['rank'] = st.groupby('om')['nn'].rank(ascending=False, method='first')
        hub_pts = set(st[(st['rank'] == 1) & (st['nn'] >= PLACEHOLDER_MIN_ROWS)
                         & (st['share'] >= PLACEHOLDER_MIN_SHARE)]['_p'])
    else:
        vc = d['_p'].value_counts()
        hub_pts = set(vc[vc >= PLACEHOLDER_ROWS_NO_OMSU].index)

    # коды: из реестра — прежние; новым — следующие номера (по ОМСУ, с севера на юг)
    used = pd.to_numeric(regp['Номер'], errors='coerce')
    next_num = int(used.max()) + 1 if len(used.dropna()) else 1
    order_new = sorted((k for k, o in enumerate(objects) if o['reg'] is None),
                       key=lambda k: (str(d.loc[d['_o'] == k, omsu_c].pipe(_mode)) if omsu_c else '',
                                      -objects[k]['lat']))
    for k, o in enumerate(objects):
        if o['reg'] is not None:
            o['code'] = str(regp.at[o['reg'], 'Код объекта'])
            o['num'] = int(regp.at[o['reg'], 'Номер'])
    for k in order_new:
        objects[k]['num'] = next_num
        objects[k]['code'] = f'{prefix}-{next_num:0{CODE_DIGITS}d}'
        next_num += 1
    for o in objects:
        o['hub'] = any(i in hub_pts for i in o['members']) or \
            (o['reg'] is not None and str(regp.at[o['reg'], 'Тип точки']) == 'условная точка округа')
        o['type'] = 'условная точка округа' if o['hub'] else 'объект'

    # 6. листы
    d['_code'] = d['_o'].map(lambda k: objects[k]['code'])
    d['_olat'] = d['_o'].map(lambda k: objects[k]['lat'])
    d['_olon'] = d['_o'].map(lambda k: objects[k]['lon'])
    d['_dist'] = np.round(haversine(d['_lat'].to_numpy(float), d['_lon'].to_numpy(float),
                                    d['_olat'].to_numpy(float), d['_olon'].to_numpy(float))).astype(int)
    cnt = d['_o'].value_counts()

    obj_rows = []
    for k in sorted(range(len(objects)), key=lambda k: objects[k]['num']):
        o = objects[k]
        if len(o['members']) > 1000:
            raise ToolError('На одном объекте больше 1 000 разных точек. Разделите выборку.')
        sub = d[d['_o'] == k]
        la, lo = P_lat[o['members']], P_lon[o['members']]
        spread = float(haversine(la[:, None], lo[:, None], la[None, :], lo[None, :]).max()) if len(la) > 1 else 0.0
        note = []
        if o['hub']:
            if omsu_c:
                om = sub[omsu_c].pipe(_mode)
                share = len(sub) / (d[omsu_c] == om).sum()
                note.append(f'условная точка округа: {len(sub)} записей ({share:.0%} записей округа) '
                            f'в одной точке — реального местоположения нет')
            else:
                note.append(f'{len(sub)} записей в одной точке — похоже на условную точку')
        if omsu_c and sub[omsu_c].nunique() > 1:
            note.append('записи разных ОМСУ')
        if rso_c and sub[rso_c].nunique() > 1:
            note.append('записи разных РСО')
        row = {'Код объекта': o['code'], 'Тип точки': o['type'], 'Широта': o['lat'], 'Долгота': o['lon'],
               'Записей': len(sub), 'Разных точек': len(o['members']), 'Разброс точек, м': round(spread)}
        if omsu_c:
            row[omsu_c] = '; '.join(str(x) for x in sub[omsu_c].dropna().unique())
        if rso_c:
            row[rso_c] = '; '.join(str(x) for x in sub[rso_c].dropna().unique())
        if id_c:
            row[f'{id_c} (список)'] = ', '.join(sorted(map(str, sub[id_c].dropna()), key=lambda x: (len(x), x)))
        row['Код из реестра'] = 'да' if o['reg'] is not None else 'новый'
        row['Примечание'] = '; '.join(note)
        obj_rows.append(row)
    objs = pd.DataFrame(obj_rows)

    rec = d[cols].copy()
    rec.insert(0, 'Код объекта', d['_code'])
    rec['Широта объекта'] = d['_olat']
    rec['Долгота объекта'] = d['_olon']
    rec['Расстояние до точки объекта, м'] = d['_dist']
    rec['Записей на объекте'] = d['_o'].map(cnt)
    rec['Тип точки'] = d['_o'].map(lambda k: objects[k]['type'])
    rec['Примечание по координатам'] = d['_note']
    rec['_num'] = d['_o'].map(lambda k: objects[k]['num'])
    try:
        rec = rec.sort_values(['_num'] + ([id_c] if id_c else []), kind='stable')
    except TypeError:                                   # ID разных типов (числа и текст)
        rec = rec.sort_values('_num', kind='stable')
    rec = rec.drop(columns='_num')

    nocoord = bad[cols].copy()
    nocoord['Причина'] = bad['_note']

    # 7. реестр
    if USE_REGISTRY:
        now = datetime.now().strftime('%Y-%m-%d %H:%M')
        add = [{'Код объекта': objects[k]['code'], 'Префикс': prefix, 'Номер': objects[k]['num'],
                'Широта': objects[k]['lat'], 'Долгота': objects[k]['lon'], 'Тип точки': objects[k]['type'],
                'Файл': name, 'Добавлен': now} for k in order_new]
        full = pd.concat([reg, pd.DataFrame(add, columns=REG_COLS)], ignore_index=True)
        # объект из реестра впервые оказался условной точкой — отметить
        for o in objects:
            if o['reg'] is not None and o['hub']:
                full.loc[full['Код объекта'] == o['code'], 'Тип точки'] = 'условная точка округа'

    if len(full) > 50000:
        raise ToolError('В реестре больше 50 000 объектов. Обратитесь к администратору.')
    output = job / 'Объекты.xlsx'
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        objs.to_excel(writer, sheet_name='Объекты', index=False)
        rec.to_excel(writer, sheet_name='Записи', index=False)
        nocoord.to_excel(writer, sheet_name='Без координат', index=False)
        full.to_excel(writer, sheet_name='Реестр', index=False)
    style(output)
    result = {'file': output.name, 'rows': len(src), 'objects': len(objects), 'unmatched': len(bad),
              'message': f'Объектов: {len(objects)}. Новых: {len(order_new)}, из реестра: {len(objects) - len(order_new)}. '
                         f'Строк без координат: {len(bad)}. Результат и реестр сохранены.'}
    return result, full


def read_sheet(data, registry=False):
    validate_office(data)
    try:
        with io.BytesIO(data) as stream:
            book = load_workbook(stream, read_only=True, data_only=True)
            try:
                sheet = book['Реестр'] if registry and 'Реестр' in book.sheetnames else book.worksheets[0]
                if (sheet.max_row or 0) > 30001 or (sheet.max_column or 0) > 200:
                    raise ToolError('Допускается до 30 000 строк и 200 столбцов.')
                rows = []
                cells = 0
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > 500000 or len(rows) > 30000:
                        raise ToolError('Слишком много данных: до 30 000 строк и 500 000 ячеек.')
                    rows.append(row)
            finally:
                book.close()
        if not rows or not rows[0]:
            raise ToolError('Файл не содержит таблицы.')
        columns = [str(value).strip() if value is not None else f'Столбец {i+1}' for i,value in enumerate(rows[0])]
        if len(columns) != len(set(columns)):
            raise ToolError('В файле повторяются заголовки столбцов.')
        return pd.DataFrame(rows[1:], columns=columns).dropna(how='all')
    except ToolError:
        raise
    except Exception:
        raise ToolError('Не удалось прочитать XLSX. Проверьте файл и заголовки первого листа.') from None


def validate_registry(frame):
    if any(c not in frame.columns for c in REG_COLS):
        raise ToolError('В реестре нет обязательных столбцов: код, префикс, номер, широта и долгота.')
    output = []
    seen = set()
    for row in frame[REG_COLS].to_dict('records'):
        try:
            code, prefix = str(row['Код объекта']), str(row['Префикс'])
            number = int(row['Номер'])
            lat, lon = float(row['Широта']), float(row['Долгота'])
            if not re.fullmatch(r'[A-ZА-ЯЁ0-9]{1,12}', prefix) or number < 1 or number != float(row['Номер']):
                raise ValueError()
            if code != f'{prefix}-{number:04d}' or code in seen or not _in_box(lat, lon):
                raise ValueError()
            seen.add(code)
            row.update({'Код объекта': code, 'Префикс': prefix, 'Номер': number, 'Широта': lat, 'Долгота': lon})
            for key in ('Тип точки', 'Файл', 'Добавлен'):
                row[key] = '' if pd.isna(row[key]) else str(row[key])[:180]
            output.append(row)
        except (ValueError, TypeError, OverflowError):
            raise ToolError('Реестр содержит повторяющиеся коды, неверные номера или координаты.') from None
    if len(output) > 50000:
        raise ToolError('В реестре больше 50 000 объектов.')
    return pd.DataFrame(output, columns=REG_COLS)


def group(name, data, radius, prefix, root, job, imported=None):
    try:
        radius = float(str(radius).replace(',', '.'))
    except ValueError:
        raise ToolError('Укажите радиус от 1 до 5 000 метров.') from None
    if not math.isfinite(radius) or not 1 <= radius <= 5000:
        raise ToolError('Укажите радиус от 1 до 5 000 метров.')
    prefix = str(prefix or '').strip().upper().rstrip('-') or prefix_for(name)
    if not re.fullmatch(r'[A-ZА-ЯЁ0-9]{1,12}', prefix):
        raise ToolError('Префикс: до 12 букв или цифр, например VZ.')
    src = read_sheet(data)
    # Keep original columns without colliding with generated output headers.
    reserved = {'Код объекта', 'Широта объекта', 'Долгота объекта', 'Расстояние до точки объекта, м',
                'Записей на объекте', 'Тип точки', 'Примечание по координатам', 'Причина'}
    rename = {}
    occupied = set(src.columns)
    for column in src.columns:
        if column in reserved or column.startswith('_'):
            target = column + ' (исходный)'
            while target in occupied:
                target += ' 2'
            rename[column] = target; occupied.add(target)
    src = src.rename(columns=rename)
    reg_path = root / 'coordinate_registry.json'
    # Serialize the read/assign/write transaction, including concurrent same-account jobs.
    with _REGISTRY_LOCK:
        try:
            registry = validate_registry(pd.DataFrame(json.loads(reg_path.read_text(encoding='utf-8')), columns=REG_COLS)) if reg_path.exists() else pd.DataFrame(columns=REG_COLS)
        except (OSError, ValueError) as exc:
            raise ToolError('Не удалось прочитать личный реестр. Данные не перезаписаны.') from exc
        if imported:
            extra = validate_registry(read_sheet(imported, registry=True))
            by_code = {row['Код объекта']: row for row in registry.to_dict('records')}
            for row in extra.to_dict('records'):
                previous = by_code.get(row['Код объекта'])
                if previous and any(previous[k] != row[k] for k in ('Префикс', 'Номер', 'Широта', 'Долгота')):
                    raise ToolError('Код из загружаемого реестра уже относится к другому объекту. Реестр не изменён.')
                if not previous:
                    by_code[row['Код объекта']] = row
            registry = pd.DataFrame(by_code.values(), columns=REG_COLS)
        result, full = _process(src, registry, name, radius, prefix, job)
        # Revalidate before committing; a failed run must not consume codes.
        full = validate_registry(full)
        tmp = reg_path.with_suffix('.tmp')
        tmp.write_text(json.dumps(full.to_dict('records'), ensure_ascii=False, allow_nan=False), encoding='utf-8')
        tmp.replace(reg_path)
        return result
