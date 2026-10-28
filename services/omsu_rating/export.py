"""Portable exports of the same calculated rating shown in the application.

Workbook columns, period selection and nonparticipants follow exportXlsx in
the supplied application. The offline snapshot contains no access credentials
or source requests; all interactions use its immutable embedded rating DTO.
"""
import base64
import hashlib
import io
import json
from html import escape
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import calculate

ROOT = Path(__file__).resolve().parents[2]
ZONES = {'z1': ('E4F2E5', '25632A'), 'z2': ('FBF0D9', '805600'), 'z3': ('FBE7E7', 'A61B1B')}


def _literal(value):
    return ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value


def _comment(cell, text):
    if text:
        cell.comment = Comment(_literal(str(text)), 'Нейрона ИИ · методика 2026')
        cell.comment.width = 480
        cell.comment.height = min(480, max(110, 20 * len(str(text).splitlines())))


def _row_values(indicator, row):
    """Original export includes all municipalities, regardless of UI filters."""
    if row.get('na'):
        return ['—', '', row['omsu'], 'ЗАТО, не участвует' if row['group'] == 'zato' else 'не участвует']
    extra = row['extra']
    extras = ([extra['fact'], ' / '.join(map(str, extra['plan']))] if indicator == 4 else
              [extra.get('without', ''), extra.get('total', '')] if indicator == 6 else
              [f"{extra['done']} / {extra['plan']}"] if indicator == 7 else [])
    names = row.get('rso_names', [detail['name'] for detail in row['details']])
    if indicator == 7:
        names = [f"{detail['name']}: {calculate.fmt(detail['value'])}" +
                 ('' if detail['found'] else ' (нет в таблице)') for detail in row['details']]
    return [row.get('place_label', str(row['place'])), row.get('delta'), row['omsu'],
            ', '.join(label for label in ['ЗАТО' if row['group'] == 'zato' else '',
                                         'обнулён' if row['no_docs'] else ''] if label),
            calculate.rnd(row['value']), *extras, ', '.join(names)]


def workbook_bytes(saved, indicator=None, period=None):
    """Four available indicators; selected period only affects its own sheet."""
    book = Workbook()
    book.remove(book.active)
    for key in (4, 5, 6, 7):
        item = saved['indicators'][str(key)]
        selected = period if indicator == key and period is not None else item['default_period']
        if selected is None or not item['periods'][selected]:
            continue
        rows = item['periods'][selected]
        extras = ['Факт, % от НВВ', 'План, %'] if key == 4 else ['ЛС без ИПУ', 'Всего ЛС'] if key == 6 else ['Мероприятий выполнено'] if key == 7 else []
        headers = ['Место', 'Динамика (+ поднялся)', 'ОМСУ', 'Группа', item['metric'], *extras, 'РСО']
        sheet = book.create_sheet(f'Показатель {key}')
        sheet.append([f"{item.get('roman', key)}. {item['name']} — {calculate.PERIODS[selected]}"])
        sheet.append(['Места — по методике расчёта на 2026 год'])
        sheet.append(['Получено: ' + (item.get('updated_at') or 'Нет данных') + ' · ' + item['note'] +
                      (' · ' + item['error'] if item.get('error') else '')])
        sheet.append(headers)
        nonparticipants = item['period_stats'][selected]['nonparticipants']
        ordered = [row for group in ('main', 'zato') for row in
                   [*filter(lambda row: row['group'] == group, rows),
                    *filter(lambda row: row['group'] == group, nonparticipants)]]
        for row in ordered:
            sheet.append([_literal(value) for value in _row_values(key, row)])
            excel_row = sheet.max_row
            if row.get('na'):
                for cell in sheet[excel_row]:
                    cell.fill = PatternFill('solid', fgColor='F1F3F6')
                    cell.font = Font(color='687181')
                _comment(sheet.cell(excel_row, 1), 'Не участвует в критерии — места нет.')
                continue
            background, foreground = ZONES[row['zone']]
            sheet.cell(excel_row, 3).fill = PatternFill('solid', fgColor=background)
            sheet.cell(excel_row, 3).font = Font(color=foreground, bold=True)
            _comment(sheet.cell(excel_row, 1), row.get('place_tooltip'))
            _comment(sheet.cell(excel_row, 3), 'Цвет: верхняя / средняя / нижняя треть полного рейтинга своей группы, до фильтрации.')
            _comment(sheet.cell(excel_row, 5), row.get('value_tooltip'))
            if row.get('delta') is not None:
                sheet.cell(excel_row, 2).number_format = '+0;-0;0'
                sheet.cell(excel_row, 2).font = Font(color='25632A' if row['delta'] > 0 else 'A61B1B' if row['delta'] < 0 else '687181')
                _comment(sheet.cell(excel_row, 2), f"{calculate.PERIODS[selected - 1]}: {row['previous_place']} место → {calculate.PERIODS[selected]}: {row['place']} место.")
            sheet.cell(excel_row, 5).number_format = '0.00'
            if key == 4:
                sheet.cell(excel_row, 6).number_format = '0.00'
                _comment(sheet.cell(excel_row, 6), row['extra'].get('fact_tooltip'))
                _comment(sheet.cell(excel_row, 7), row['extra'].get('plan_tooltip'))
            elif key == 6:
                _comment(sheet.cell(excel_row, 6), row['extra'].get('without_tooltip'))
            if row['no_docs']:
                sheet.cell(excel_row, 4).font = Font(color='A61B1B', bold=True)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _literal(cell.value)
                    cell.data_type = 's'  # A source name beginning '=' is never an Excel formula.
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[4]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='515DCE')
        widths = [11, 19, 28, 20, 24, *([24] * len(extras)), 58]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for title_row in (1, 2, 3):
            sheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=len(headers))
        sheet.row_dimensions[1].height = 30
        sheet.row_dimensions[2].height = 24
        sheet.row_dimensions[3].height = 56
        sheet.row_dimensions[4].height = 34
        sheet.freeze_panes = 'D5'
        sheet.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{sheet.max_row}'
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:4'
    if not book.worksheets:
        sheet = book.create_sheet('Нет данных')
        sheet.append(['Рейтинг ОМСУ · Нейрона ИИ'])
        sheet.append(['Пока нет данных. Получите данные источников в Нейроне и сохраните выгрузку повторно.'])
        sheet.column_dimensions['A'].width = 90
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


SNAPSHOT_SHELL = """
:root{color-scheme:light;--text:#1b293d;--muted:#59677c;--card:#ffffffeb;--border:#d7deec;--primary:#6254df;--bg-soft:#f4f6fb;--violet-soft:#efedff}
*{box-sizing:border-box}body{margin:0;font:15px/1.5 system-ui,-apple-system,sans-serif;color:var(--text);background:linear-gradient(130deg,#e6edff,#f6f2ff);min-height:100vh}.container{max-width:1540px;padding:24px;margin:auto}.topbar{display:flex;justify-content:space-between;padding:18px 24px;border:1px solid var(--border);border-radius:22px;background:var(--card)}.brand{display:flex;align-items:center;gap:12px}.brand-logo{width:40px;height:40px}.brand-title{font-size:18px;font-weight:750}.brand-subtitle{font-size:12px;color:var(--muted)}[hidden]{display:none!important}button,input,select{font:inherit}@media(max-width:700px){.container{padding:12px}.topbar{padding:14px}}@media print{.or-actions,.or-tabs,.or-filters,.top-actions{display:none!important}.container{padding:0}.or-table-wrap{overflow:visible}.or-panel,.or-hero{box-shadow:none;border:0}.or-table{min-width:0!important;font-size:9px}.or-table td,.or-table th{padding:5px}.or-detail{position:static;max-height:none;width:100%}}
"""


def snapshot_html(saved):
    """Inline the native UI with a CSP that permits only its exact script hash."""
    # Explicitly copy only the module response contract, never session state or raw issues.
    portable = {key: saved[key] for key in ('methodology_year', 'period_labels', 'municipalities', 'indicators', 'updated_at', 'refresh', 'history')}
    portable['offline_exports'] = {'default': base64.b64encode(workbook_bytes(saved)).decode('ascii')}
    for indicator, item in saved['indicators'].items():
        for period, rows in enumerate(item['periods']):
            if rows:
                portable['offline_exports'][f'{indicator}:{period}'] = base64.b64encode(workbook_bytes(saved, int(indicator), period)).decode('ascii')
    payload = json.dumps(portable, ensure_ascii=False, separators=(',', ':'), allow_nan=False)
    payload = payload.replace('&', '\\u0026').replace('<', '\\u003c').replace('>', '\\u003e').replace('\u2028', '\\u2028').replace('\u2029', '\\u2029')
    script = (ROOT / 'static/omsu-rating.js').read_text(encoding='utf-8')
    stylesheet = (ROOT / 'static/omsu-rating.css').read_text(encoding='utf-8')
    digest = base64.b64encode(hashlib.sha256(script.encode('utf-8')).digest()).decode('ascii')
    policy = f"default-src 'none'; script-src 'sha256-{digest}'; style-src 'unsafe-inline'; img-src data:; connect-src 'none'; object-src 'none'; base-uri 'none'; form-action 'none'"
    environment = Environment(loader=FileSystemLoader(ROOT / 'templates'), autoescape=select_autoescape(['html']))
    content = environment.get_template('omsu_rating_content.html').render(snapshot=True)
    logo = base64.b64encode((ROOT / 'static/logo.svg').read_bytes()).decode('ascii')
    content = content.replace('src="/static/logo.svg"', f'src="data:image/svg+xml;base64,{logo}"')
    return ('<!doctype html><html lang="ru" data-theme="light"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width,initial-scale=1">'
            '<meta http-equiv="Content-Security-Policy" content="' + escape(policy, quote=True) + '">'
            '<title>Рейтинг ОМСУ — сохранённый снимок · Нейрона ИИ</title><style>' + SNAPSHOT_SHELL + stylesheet + '</style>'
            '</head><body class="unified-module-page module-omsu-rating">' + content +
            '<noscript>Для просмотра вкладок и расчётов в сохранённом снимке включите JavaScript. Подключение к сети не требуется.</noscript>'
            '<script id="or-snapshot-data" type="application/json">' + payload + '</script><script>' + script + '</script></body></html>')
