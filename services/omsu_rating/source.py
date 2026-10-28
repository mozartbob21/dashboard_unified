"""Fixed, read-only sources; upstream response bodies and credentials never reach clients."""
import io
import time
import zipfile
from datetime import datetime, date, time as clock_time, timedelta
from urllib.parse import urlsplit
import httpx
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
from .methodology import REDMINE_ORIGIN, GSHEET_ID

class SourceError(Exception): pass
MAX_BYTES=20*1024*1024
FIELDS={17,70,158,159,160,161,165,166,167,168,169,208,209,210,211,212,213,259,260,261,*range(184,208)}

def limited(client,url,*,headers=None,params=None):
    with client.stream('GET',url,headers=headers,params=params) as response:
        if response.status_code in (301,302,303,307,308):return None,response.headers.get('location','')
        if response.status_code in (401,403):raise SourceError('Источник отклонил доступ. Проверьте API-ключ и права в системе-источнике.')
        if response.status_code!=200:raise SourceError('Источник временно недоступен. Сохранённые данные не изменены.')
        body=bytearray()
        for block in response.iter_bytes():
            body.extend(block)
            if len(body)>MAX_BYTES:raise SourceError('Ответ источника превышает допустимый размер.')
        return bytes(body),None

def fetch_redmine(query,key,client=None):
    import json
    own=client is None
    client=client or httpx.Client(timeout=httpx.Timeout(25,connect=5),verify=True,follow_redirects=False,trust_env=False)
    try:
        all_rows=[];seen=set();expected_total=None;offset=0;deadline=time.monotonic()+150
        while offset<10000:
            if time.monotonic()>deadline:raise SourceError('Источник не успел передать все данные. Повторите обновление позднее.')
            # The saved query defines ordering. Overriding it changes the original
            # stable order of municipalities with equal indicator values.
            body,redirect=limited(client,REDMINE_ORIGIN+'/issues.json',headers={'X-Redmine-API-Key':key},params={'query_id':query,'status_id':'*','limit':100,'offset':offset})
            if redirect is not None:raise SourceError('Источник перенаправил запрос. Проверьте доступ по API.')
            try:
                result=json.loads(body);rows=result['issues'];total=result['total_count']
                if not isinstance(rows,list) or type(total) is not int or total<0 or total>10000 or len(rows)>100 or offset+len(rows)>total:raise ValueError()
                if expected_total is not None and total!=expected_total:raise SourceError('Список задач изменился во время загрузки. Повторите обновление; прежние данные сохранены.')
                expected_total=total
                for issue in rows:
                    if not isinstance(issue,dict) or not isinstance(issue.get('custom_fields',[]),list):raise ValueError()
                    ident=issue.get('id')
                    if type(ident) is not int or ident<=0 or ident in seen:raise SourceError('Источник вернул повторяющиеся или некорректные задачи. Прежние данные сохранены.')
                    seen.add(ident)
                    assigned=issue.get('assigned_to') or {}
                    all_rows.append({'id':issue.get('id'),'subject':str(issue.get('subject') or '')[:1000],
                        'assigned_to':{'name':str(assigned.get('name') or '')[:200]},
                        'custom_fields':[{'id':f['id'],'value':f.get('value')} for f in issue.get('custom_fields',[]) if isinstance(f,dict) and f.get('id') in FIELDS]})
            except (ValueError,KeyError,TypeError,AttributeError):raise SourceError('Источник вернул данные неизвестного формата.') from None
            if offset+len(rows)>=total:return all_rows
            if not rows:raise SourceError('Источник вернул неполный список задач. Старые данные сохранены.')
            offset+=len(rows)
        raise SourceError('Слишком много задач для одной проверки.')
    except httpx.HTTPError:raise SourceError('Не удалось установить защищённое соединение с источником.') from None
    finally:
        if own:client.close()

def fetch_sheet(client=None):
    own=client is None
    client=client or httpx.Client(timeout=httpx.Timeout(30,connect=5),verify=True,follow_redirects=False,trust_env=False)
    url=f'https://docs.google.com/spreadsheets/d/{GSHEET_ID}/export?format=xlsx'
    try:
        for _ in range(4):
            body,redirect=limited(client,url)
            if redirect is None:break
            target=urlsplit(redirect)
            if target.scheme!='https' or target.username or target.password or target.port not in (None,443) or not (target.hostname=='docs.google.com' or (target.hostname or '').endswith('.googleusercontent.com')):
                raise SourceError('Источник таблицы перенаправил запрос на неподтверждённый адрес.')
            url=redirect
        else:raise SourceError('Источник таблицы не завершил перенаправление.')
        with zipfile.ZipFile(io.BytesIO(body)) as z:
            if sum(i.file_size for i in z.infolist())>100*1024*1024:raise SourceError('Таблица слишком велика для обработки.')
        book=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
        try:
            name=next((n for n in book.sheetnames if n.startswith('Исполнение дорожной')),None)
            if name is None:raise SourceError('В источнике нет листа исполнения дорожной карты.')
            rows=[]
            for row in book[name].iter_rows(min_row=2,max_col=22,values_only=True):
                if len(rows)>=10000:raise SourceError('Слишком много строк в таблице дорожных карт.')
                # K7 recognises phrases anywhere in a response. Truncating cell
                # text before classification silently changes the methodology.
                # Compressed and expanded workbook size limits apply above.
                # The EXE's online interface uses SheetJS raw cells (cellDates
                # is false): date-formatted cells remain Excel serial numbers.
                # openpyxl decodes them eagerly, so restore their raw numeric
                # meaning before applying the original K7 text/number rules.
                rows.append([to_excel(v,book.epoch) if isinstance(v,(datetime,date,clock_time,timedelta)) else
                             v if isinstance(v,(int,float,bool)) or v is None else str(v) for v in row])
            return rows
        finally:book.close()
    except SourceError:raise
    except Exception:raise SourceError('Не удалось получить таблицу дорожных карт. Показатели 4–6 доступны отдельно.') from None
    finally:
        if own:client.close()
